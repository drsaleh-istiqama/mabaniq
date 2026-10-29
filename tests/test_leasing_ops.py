"""0.22.0 — leasing operations: renewal, termination, eviction, hand-over and deposit settlement, arrears notices, payment plans, late-payment charity,
maintenance with tenant charge-back, owner statements (the formula), the scoped property manager (leak test) and the Excel/CSV import (dry-run / commit / re-run)."""
import datetime as dt
import io
import os
import tempfile

os.environ.setdefault("MABANIQ_DB", os.path.join(tempfile.mkdtemp(), "leaops.db"))
os.environ["MABANIQ_FORCE_PW_CHANGE"] = "0"
os.environ["MABANIQ_TENANTS"] = "bayan:بيان عمان;almaqsura:المقصورة العقارية"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from backend import auth as A  # noqa: E402
from backend import leasing as L  # noqa: E402
from backend.app import app  # noqa: E402
from backend.db import TENANT, connect  # noqa: E402

CSRF = "csrf-leaops-" + "o" * 28
D = dt.date
TODAY = dt.date.today()
_N = [0]
_UNITS = "SELECT id FROM units WHERE code LIKE 'T-OPS-%'"
_LEASES = f"SELECT id FROM leases WHERE unit_id IN ({_UNITS})"


def login_as(c, user, tenant="bayan"):
    tok = TENANT.set(tenant)
    try:
        c.cookies.set(A.COOKIE, A.issue_session(user))
    finally:
        TENANT.reset(tok)
    c.cookies.set(A.TENANT_COOKIE, tenant)
    c.cookies.set(A.CSRF_COOKIE, CSRF)
    c.headers["X-CSRF-Token"] = CSRF
    return c


@pytest.fixture()
def cl():
    from backend import observability as O
    O.limiter.reset()
    with TestClient(app, base_url="https://testserver") as c:
        for t in ("jadwa", "bayan", "almaqsura"):
            tok = TENANT.set(t)
            try:
                d = connect()
                d.execute("DELETE FROM login_attempts")
                d.commit()
                d.close()
            finally:
                TENANT.reset(tok)
        yield c


def _sql(sql, args=(), commit=True):
    """Direct access for the few things the API deliberately cannot do (move a date into the past, read a table)."""
    tok = TENANT.set("bayan")
    try:
        c = connect()
        r = c.execute(sql, args)
        out = r.fetchall() if sql.lstrip().upper().startswith("SELECT") else None
        if commit:
            c.commit()
        c.close()
        return out
    finally:
        TENANT.reset(tok)


@pytest.fixture(scope="module", autouse=True)
def _purge_my_rows():
    yield
    tok = TENANT.set("bayan")
    try:
        c = connect()
        pg = getattr(c, "dialect", "sqlite") == "postgres"
        if pg:
            c.maintenance()
        sc = " AND org_id = current_setting('app.org_id', true)" if pg else ""
        stmts = [f"DELETE FROM meter_readings WHERE handover_id IN (SELECT id FROM lease_handovers WHERE lease_id IN ({_LEASES}{sc}))",
                 f"DELETE FROM lease_handovers WHERE lease_id IN ({_LEASES}{sc})", f"DELETE FROM arrears_notices WHERE lease_id IN ({_LEASES}{sc})",
                 f"DELETE FROM payment_plans WHERE lease_id IN ({_LEASES}{sc})", f"DELETE FROM charity_dues WHERE lease_id IN ({_LEASES}{sc})",
                 f"DELETE FROM maint_orders WHERE request_id IN (SELECT id FROM maint_requests WHERE unit_id IN ({_UNITS}{sc}))",
                 f"DELETE FROM maint_requests WHERE unit_id IN ({_UNITS}{sc})", f"DELETE FROM property_expenses WHERE unit_id IN ({_UNITS}{sc})",
                 f"DELETE FROM owner_statements WHERE owner_id IN (SELECT id FROM owners WHERE name LIKE 'مالك اختبار%'{sc})",
                 f"DELETE FROM user_scopes WHERE user_id IN (SELECT id FROM users WHERE username LIKE 'pmt-%'{sc})",
                 f"DELETE FROM sessions WHERE user_id IN (SELECT id FROM users WHERE username LIKE 'pmt-%'{sc})", f"DELETE FROM users WHERE username LIKE 'pmt-%'{sc}",
                 f"DELETE FROM rent_payments WHERE lease_id IN ({_LEASES}{sc})", f"DELETE FROM cheques WHERE lease_id IN ({_LEASES}{sc})",
                 f"DELETE FROM lease_events WHERE lease_id IN ({_LEASES}{sc})", f"DELETE FROM lease_terms WHERE lease_id IN ({_LEASES}{sc})",
                 f"DELETE FROM rent_dues WHERE lease_id IN ({_LEASES}{sc})",
                 f"DELETE FROM tenants_l WHERE id IN (SELECT tenant_id FROM leases WHERE unit_id IN ({_UNITS}{sc})) OR name LIKE 'مستأجر استيراد%'{sc}",
                 f"DELETE FROM leases WHERE unit_id IN ({_UNITS}{sc})", f"DELETE FROM unit_meters WHERE unit_id IN ({_UNITS}{sc})",
                 f"DELETE FROM units WHERE code LIKE 'T-OPS-%'{sc}", f"DELETE FROM owners WHERE name LIKE 'مالك اختبار%'{sc}", f"DELETE FROM vendors WHERE name LIKE 'مورّد اختبار%'{sc}",
                 f"DELETE FROM import_batches WHERE filename LIKE 'ops-test-%'{sc}",
                 f"DELETE FROM doc_issues WHERE kind IN ('lease','rent_receipt','cheque_receipt','handover_report','arrears_notice','owner_statement'){sc}"]
        for sql in stmts:
            c.execute(sql)
        c.commit()
        if pg:
            c.end_maintenance()
        c.close()
    finally:
        TENANT.reset(tok)


def _unit(cl, owner_id=None):
    tok = TENANT.set("bayan")
    try:
        c = connect()
        _N[0] += 1
        row = c.execute("SELECT project_id FROM units WHERE retained=1 ORDER BY project_id LIMIT 1").fetchone()  # the project the portfolio screen lists first
        uid = c.execute("INSERT INTO units(project_id,code,type,area,status,retained,use,owner_id) VALUES(?,?,?,?,?,?,?,?)",
                        (row[0], f"T-OPS-{_N[0]:03d}", "مكتب", 60, "a", 1, "commercial", owner_id)).lastrowid
        c.commit()
        c.close()
    finally:
        TENANT.reset(tok)
    return uid


def _lease(cl, uid=None, **kw):
    body = {"unit_id": uid or _unit(cl), "tenant": {"name": "شركة الرواد للتجارة", "phone": "+968 91234567", "id_number": "OP123456"},
            "start": "2035-01-01", "months": 12, "annual_rent": 1200, "frequency": 4, "sign": True} | kw
    r = cl.post("/api/lea/leases", json=body)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _detail(cl, lid):
    return cl.get(f"/api/lea/leases/{lid}").json()


def _past_start(days):
    return (TODAY - dt.timedelta(days=days)).isoformat()


def _no_mismatch():
    tok = TENANT.set("bayan")
    try:
        c = connect()
        bad = L.ledger_mismatches(c)
        c.close()
    finally:
        TENANT.reset(tok)
    assert bad == []


# ----------------------------------------------------------------------------------------------------- renewal
def test_renew_continues_the_lease_and_carries_the_deposit(cl):
    login_as(cl, "admin")
    lid = _lease(cl, deposit=200)
    dep = next(d for d in _detail(cl, lid)["dues"] if d["kind"] == "deposit")
    assert cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 200, "due_id": dep["id"], "method": "cash"}).status_code == 200
    r = cl.post(f"/api/lea/leases/{lid}/renew", json={"months": 12, "annual_rent": 1300})
    assert r.status_code == 200 and r.json()["status"] == "active", r.text
    new, old = _detail(cl, r.json()["id"]), _detail(cl, lid)
    assert new["start"] == "2036-01-01" and new["renewed_from"] == lid and new["deposit_state"] == "carried" and not [d for d in new["dues"] if d["kind"] == "deposit"]
    assert old["status"] == "renewed" and old["deposit_state"] == "transferred"
    assert cl.post(f"/api/lea/leases/{lid}/renew", json={"months": 12}).status_code == 409  # already renewed
    assert cl.post(f"/api/lea/leases/{lid}/deposit/settle", json={}).status_code == 409  # the deposit moved with the tenant
    ops = cl.get(f"/api/lea/leases/{lid}/ops").json()
    assert ops["renewed_to"]["id"] == new["id"]
    _no_mismatch()


# ----------------------------------------------------------------------------------------------------- termination
def test_terminate_cuts_future_dues_but_keeps_what_was_paid(cl):
    login_as(cl, "admin")
    lid = _lease(cl)
    dues = _detail(cl, lid)["dues"]
    assert [d["amount"] for d in dues] == [300, 300, 300, 300]
    cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 300, "due_id": dues[0]["id"], "method": "cash"})
    cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 100, "due_id": dues[1]["id"], "method": "cash"})  # partly paid, but its date is after the exit
    ch = cl.post("/api/lea/cheques", json={"lease_id": lid, "due_id": dues[3]["id"], "number": "T-OPS-1", "bank": "بنك مسقط", "amount": 300, "due_date": "2035-10-01"}).json()
    r = cl.post(f"/api/lea/leases/{lid}/terminate", json={"on": "2035-03-15", "note": "إنهاء بالتراضي", "fee": 50})
    assert r.status_code == 200, r.text
    assert r.json()["cancelled_dues"] == 3 and r.json()["cancelled_amount"] == 800
    d = _detail(cl, lid)
    assert d["status"] == "terminated" and d["termination_kind"] == "early" and d["terminated_at"] == "2035-03-15"
    rent = [x for x in d["dues"] if x["kind"] == "rent"]
    assert [(x["amount"], x["paid"]) for x in rent if x["amount"] > 0] == [(300, 300), (100, 100)]  # the part-paid one is cut to what was paid
    assert len([x for x in rent if "أُلغيت" in (x["label"] or "")]) == 3  # the later dues stay as labelled zero lines (rent_dues has no DELETE for the app role)
    assert [x["amount"] for x in d["dues"] if x["kind"] == "fee"] == [50]
    assert [q["status"] for q in d["cheques"] if q["id"] == ch["id"]] == ["cancelled"]
    units = {u["id"]: u for u in cl.get("/api/lea/units").json()["units"]}
    assert units[d["unit_id"]]["state"] == "vacant"
    assert cl.post(f"/api/lea/leases/{lid}/terminate", json={"note": "مرة ثانية"}).status_code == 409
    _no_mismatch()


def test_terminate_can_charge_only_the_days_used_and_a_fee_needs_the_manager(cl):
    login_as(cl, "admin")
    lid = _lease(cl, start="2036-01-01", frequency=1)
    r = cl.post(f"/api/lea/leases/{lid}/terminate", json={"on": "2036-04-09", "note": "إنهاء مبكر", "prorate_last": True})
    assert r.status_code == 200
    d = _detail(cl, lid)
    assert d["dues"][0]["amount"] == round(1200 * 100 / 366, 3)  # 100 of the 366 days of leap-year 2036
    lid2 = _lease_as_admin(cl)
    login_as(cl, "finance")
    assert cl.post(f"/api/lea/leases/{lid2}/terminate", json={"on": "2037-02-01", "note": "إنهاء مبكر", "fee": 20}).status_code == 403
    assert cl.post(f"/api/lea/leases/{lid2}/terminate", json={"on": "2037-02-01", "note": "إنهاء بلا مبلغ"}).status_code == 200


def _lease_as_admin(cl, **kw):
    login_as(cl, "admin")
    return _lease(cl, start="2037-01-01", **kw)


# ----------------------------------------------------------------------------------------------------- notices + eviction
def test_notices_escalate_and_eviction_needs_an_expired_final_notice_and_the_manager(cl):
    login_as(cl, "admin")
    lid = _lease(cl, start=_past_start(60), frequency=4)
    assert cl.post(f"/api/lea/leases/{_lease(cl, start='2038-01-01')}/notices", json={"kind": "reminder"}).status_code == 409  # nothing late
    assert cl.post(f"/api/lea/leases/{lid}/notices", json={"kind": "final"}).status_code == 409  # a formal notice must come first
    n1 = cl.post(f"/api/lea/leases/{lid}/notices", json={"kind": "formal", "cure_days": 10})
    assert n1.status_code == 200 and n1.json()["number"].startswith(f"AN-{TODAY.year}-") and n1.json()["amount"] == 300
    n2 = cl.post(f"/api/lea/leases/{lid}/notices", json={"kind": "final"}).json()
    assert n2["number"] != n1.json()["number"]
    pdf = cl.get(f"/api/docs/arrears_notice/{n2['id']}.pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    assert cl.post(f"/api/lea/notices/{n2['id']}/sent", json={"via": "registered"}).status_code == 200
    assert cl.post(f"/api/lea/notices/{n2['id']}/sent", json={"via": "hand"}).status_code == 409
    assert cl.post(f"/api/lea/leases/{lid}/evict", json={"note": "لم يسدد بعد الإنذار النهائي"}).status_code == 409  # the cure period is still running
    _sql("UPDATE arrears_notices SET cure_by='2000-01-01' WHERE id=?", (n2["id"],))
    login_as(cl, "finance")
    assert cl.post(f"/api/lea/leases/{lid}/evict", json={"note": "لم يسدد بعد الإنذار النهائي"}).status_code == 403
    login_as(cl, "admin")
    r = cl.post(f"/api/lea/leases/{lid}/evict", json={"note": "لم يسدد بعد الإنذار النهائي"})
    assert r.status_code == 200 and _detail(cl, lid)["status"] == "evicted"
    assert any(n["state"] in ("expired", "running") for n in cl.get("/api/lea/notices").json())


# ----------------------------------------------------------------------------------------------------- hand-over + deposit
def test_handover_readings_and_deposit_settlement_by_documented_deductions(cl):
    login_as(cl, "admin")
    uid = _unit(cl)
    mid = cl.post(f"/api/lea/units/{uid}/meters", json={"kind": "electricity", "number": "E-1"}).json()["id"]
    lid = _lease(cl, uid, deposit=200)
    dep = next(d for d in _detail(cl, lid)["dues"] if d["kind"] == "deposit")
    cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 200, "due_id": dep["id"], "method": "transfer"})
    assert cl.post(f"/api/lea/leases/{lid}/handover", json={"kind": "in", "keys": 3, "condition": {"الجدران": "سليمة", "المكيفات": "تعمل"}, "readings": [{"meter_id": mid, "reading": 1000}]}).status_code == 200
    assert cl.post(f"/api/lea/leases/{lid}/handover", json={"kind": "in"}).status_code == 409  # one report per kind
    assert cl.post(f"/api/lea/leases/{lid}/handover", json={"kind": "out"}).status_code == 409  # the lease is still running
    assert cl.post(f"/api/lea/leases/{lid}/deposit/settle", json={}).status_code == 409
    cl.post(f"/api/lea/leases/{lid}/terminate", json={"on": "2035-06-01", "note": "إنهاء بالتراضي"})
    assert cl.post(f"/api/lea/leases/{lid}/handover", json={"kind": "out", "readings": [{"meter_id": mid, "reading": 900}]}).status_code == 400  # reads lower than at hand-over
    ho = cl.post(f"/api/lea/leases/{lid}/handover", json={"kind": "out", "keys": 3, "readings": [{"meter_id": mid, "reading": 1250}],
                                                          "deductions": [{"label": "إصلاح باب", "amount": 50, "ref": "فاتورة 114"}]})
    assert ho.status_code == 200, ho.text
    ops = cl.get(f"/api/lea/leases/{lid}/ops").json()
    out = next(h for h in ops["handovers"] if h["kind"] == "out")
    assert out["readings"][0]["consumption"] == 250 and ops["deposit"]["basis"] == 200
    pdf = cl.get(f"/api/docs/handover_report/{ho.json()['id']}.pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    r = cl.post(f"/api/lea/leases/{lid}/deposit/settle", json={}).json()  # uses the deductions written on the take-back report
    assert (r["basis"], r["deducted"], r["refund"], r["state"]) == (200, 50, 150, "partly_settled")
    assert cl.post(f"/api/lea/leases/{lid}/deposit/settle", json={}).status_code == 409  # only once


def test_deposit_deductions_cannot_exceed_it_and_can_settle_an_unpaid_due(cl):
    login_as(cl, "admin")
    lid = _lease(cl, deposit=100, start=_past_start(60), frequency=4)
    dues = _detail(cl, lid)["dues"]
    dep = next(d for d in dues if d["kind"] == "deposit")
    late = next(d for d in dues if d["kind"] == "rent")
    cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 100, "due_id": dep["id"], "method": "cash"})
    cl.post(f"/api/lea/leases/{lid}/terminate", json={"on": TODAY.isoformat(), "note": "إنهاء مبكر"})
    big = {"deductions": [{"label": "تلفيات", "amount": 150, "ref": "تقرير 7"}]}
    assert cl.post(f"/api/lea/leases/{lid}/deposit/settle", json=big).status_code == 409
    r = cl.post(f"/api/lea/leases/{lid}/deposit/settle", json={"deductions": [{"label": "إيجار متأخر", "amount": 100, "ref": "اتفاق 3", "due_id": late["id"]}]})
    assert r.status_code == 200 and r.json()["state"] == "settled" and r.json()["refund"] == 0
    after = next(d for d in _detail(cl, lid)["dues"] if d["id"] == late["id"])
    assert after["paid"] == 100 and after["payments"][-1]["method"] == "other"
    _no_mismatch()


# ----------------------------------------------------------------------------------------------------- payment plans
def test_payment_plan_respreads_the_same_amount_and_only_a_manager_may_make_it(cl):
    login_as(cl, "admin")
    lid = _lease(cl, start=_past_start(200), frequency=4)
    late = [d for d in _detail(cl, lid)["dues"] if d["kind"] == "rent" and d["state"] == "late"]
    assert len(late) >= 2
    pick = late[:2]
    ch = cl.post("/api/lea/cheques", json={"lease_id": lid, "due_id": pick[0]["id"], "number": "T-OPS-P", "bank": "بنك ظفار", "amount": 300, "due_date": TODAY.isoformat()}).json()
    soon = (TODAY + dt.timedelta(days=10)).isoformat()
    later = (TODAY + dt.timedelta(days=40)).isoformat()
    body = {"due_ids": [d["id"] for d in pick], "installments": [{"due": soon, "amount": 350}, {"due": later, "amount": 250}]}
    login_as(cl, "finance")
    assert cl.post(f"/api/lea/leases/{lid}/plans", json=body).status_code == 403
    login_as(cl, "admin")
    bad = {"due_ids": [d["id"] for d in pick], "installments": [{"due": soon, "amount": 350}, {"due": later, "amount": 260}]}
    assert cl.post(f"/api/lea/leases/{lid}/plans", json=bad).status_code == 409  # a plan can't add 10 of "fee"
    before = _detail(cl, lid)["totals"]["balance"]
    ok = {"due_ids": [d["id"] for d in pick], "installments": [{"due": soon, "amount": 300}, {"due": later, "amount": 300}]}
    r = cl.post(f"/api/lea/leases/{lid}/plans", json=ok)
    assert r.status_code == 200, r.text
    d = _detail(cl, lid)
    assert d["totals"]["balance"] == before  # nothing forgiven, nothing added
    assert len([x for x in d["dues"] if x["label"] and "خطة سداد" in x["label"]]) == 2
    assert next(q for q in d["cheques"] if q["id"] == ch["id"])["due_id"] is None  # a loose cheque now clears against the oldest due
    plan = next(p for p in cl.get("/api/lea/plans").json() if p["id"] == r.json()["id"])
    assert plan["state"] == "active" and plan["total"] == 600
    cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 600, "method": "cash"})
    plan = next(p for p in cl.get("/api/lea/plans").json() if p["id"] == r.json()["id"])
    assert plan["state"] == "active" and plan["paid"] == 300  # the 600 cleared the oldest late due (300) and the first instalment (300); no interest was ever added
    _no_mismatch()


# ----------------------------------------------------------------------------------------------------- charity on late rent
def test_late_payment_charity_only_when_the_lease_has_the_clause_and_never_as_income(cl):
    login_as(cl, "admin")
    with_clause = _lease(cl, start=_past_start(90), frequency=1, charity_clause=True)
    without = _lease(cl, start=_past_start(90), frequency=1)
    due_a = _detail(cl, with_clause)["dues"][0]
    due_b = _detail(cl, without)["dues"][0]
    tok = TENANT.set("bayan")
    try:
        c = connect()
        from backend.db import setting_f
        rate, grace = setting_f(c, "late_penalty_monthly", .01), int(setting_f(c, "late_grace_days", 15))
        c.close()
    finally:
        TENANT.reset(tok)
    pay = cl.post("/api/lea/payments", json={"lease_id": with_clause, "amount": 1200, "due_id": due_a["id"], "method": "cash"}).json()
    cl.post("/api/lea/payments", json={"lease_id": without, "amount": 1200, "due_id": due_b["id"], "method": "cash"})
    items = cl.get("/api/lea/charity").json()["items"]
    mine = [i for i in items if i["lease_id"] == with_clause]
    expected = round(1200 * rate * (90 - grace) / 30, 3)
    assert len(mine) == 1 and mine[0]["amount"] == expected and not [i for i in items if i["lease_id"] == without]
    assert _detail(cl, with_clause)["totals"]["paid"] == 1200  # rent is rent: the donation is on top and never part of it
    cl.post(f"/api/lea/payments/{pay['payment_ids'][0]}/reverse", json={"reason": "شيك مرتجع"})
    again = [i for i in cl.get("/api/lea/charity").json()["items"] if i["lease_id"] == with_clause]
    assert again[0]["amount"] == 0  # taking the payment back takes the donation back
    cl.post("/api/lea/payments", json={"lease_id": with_clause, "amount": 1200, "due_id": due_a["id"], "method": "cash"})
    live = next(i for i in cl.get("/api/lea/charity").json()["items"] if i["lease_id"] == with_clause)
    assert live["amount"] == expected
    assert cl.post(f"/api/lea/charity/{live['id']}/disburse", json={"reference": "x1", "beneficiary": "جمعية"}).status_code == 409  # collect first
    assert cl.post(f"/api/lea/charity/{live['id']}/collect", json={"reference": "إيصال تبرع 1"}).status_code == 200
    assert cl.post(f"/api/lea/charity/{live['id']}/disburse", json={"reference": "تحويل 9"}).status_code == 400  # the beneficiary is required
    assert cl.post(f"/api/lea/charity/{live['id']}/disburse", json={"reference": "تحويل 9", "beneficiary": "جمعية الاستقامة"}).status_code == 200


# ----------------------------------------------------------------------------------------------------- maintenance
def test_maintenance_flow_and_tenant_charge_back(cl):
    login_as(cl, "admin")
    uid = _unit(cl)
    lid = _lease(cl, uid, start=_past_start(30))
    vid = cl.post("/api/lea/vendors", json={"name": "مورّد اختبار للتكييف", "trade": "تكييف", "iban": "OM12345678901234567"}).json()["id"]
    assert next(v for v in cl.get("/api/lea/vendors").json() if v["id"] == vid)["iban"].startswith("•")
    m = cl.post("/api/lea/maint", json={"unit_id": uid, "category": "تكييف", "description": "المكيف لا يبرّد", "priority": "urgent", "source": "tenant"})
    assert m.status_code == 200 and m.json()["sla_due"] == (TODAY + dt.timedelta(days=1)).isoformat()
    mid = m.json()["id"]
    assert cl.post(f"/api/lea/maint/{mid}/orders", json={"vendor_id": vid, "cost": 40, "charge_to": "tenant"}).status_code == 400  # a charge-back needs its basis
    o = cl.post(f"/api/lea/maint/{mid}/orders", json={"vendor_id": vid, "cost": 40, "charge_to": "tenant", "basis": "البند 6 من العقد — سوء الاستعمال بمحضر المعاينة"})
    assert o.status_code == 200
    assert cl.post(f"/api/lea/maint/{mid}/status", json={"status": "done"}).status_code == 409  # complete the order first
    login_as(cl, "finance")
    assert cl.post(f"/api/lea/maint/orders/{o.json()['id']}/done", json={"cost": 45}).status_code == 403  # billing the tenant is the manager's call
    login_as(cl, "admin")
    r = cl.post(f"/api/lea/maint/orders/{o.json()['id']}/done", json={"cost": 45})
    assert r.status_code == 200 and r.json()["tenant_due_id"]
    d = _detail(cl, lid)
    svc = [x for x in d["dues"] if x["kind"] == "service"]
    assert len(svc) == 1 and svc[0]["amount"] == 45
    exp = cl.get("/api/lea/expenses").json()
    mine = [e for e in exp if e["unit"] and e["unit"].startswith("T-OPS-") and e["order_id"] == o.json()["id"]]
    assert len(mine) == 1 and mine[0]["charge_tenant"] == 1 and mine[0]["amount"] == 45
    assert cl.get("/api/lea/maint", params={"status": "done"}).json()[0]["status"] == "done"
    assert cl.post(f"/api/lea/maint/orders/{o.json()['id']}/done", json={}).status_code == 409


# ----------------------------------------------------------------------------------------------------- owner statement
def test_owner_statement_formula_and_adjustment(cl):
    login_as(cl, "admin")
    oid = cl.post("/api/lea/owners", json={"name": "مالك اختبار الأول", "kind": "third_party", "mgmt_fee_pct": 5, "mgmt_basis": "collected"}).json()["id"]
    uid = _unit(cl, oid)
    lid = _lease(cl, uid, start="2039-01-01", frequency=1)
    p1 = cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 1000, "method": "cash"}).json()
    period = TODAY.isoformat()[:7]
    cl.post("/api/lea/expenses", json={"unit_id": uid, "category": "نظافة", "amount": 100, "charge_owner": True})
    cl.post("/api/lea/expenses", json={"unit_id": uid, "category": "ضمان مصنع", "amount": 50, "charge_owner": False})  # the company's own cost
    pv = cl.get(f"/api/lea/owners/{oid}/statement", params={"period": period}).json()
    assert (pv["collected"], pv["expenses"], pv["commission"], pv["net"]) == (1000, 100, 50, 850)  # collected − expenses − 5% of collected
    assert cl.get(f"/api/lea/owners/{oid}/statement", params={"period": "2999-01"}).status_code == 400
    s = cl.post(f"/api/lea/owners/{oid}/statements", json={"period": period})
    assert s.status_code == 200 and s.json()["net"] == 850 and s.json()["kind"] == "statement" and s.json()["number"].startswith("OS-")
    assert cl.post(f"/api/lea/owners/{oid}/statements", json={"period": period}).status_code == 409  # nothing changed since
    cl.post(f"/api/lea/payments/{p1['payment_ids'][0]}/reverse", json={"reason": "شيك مرتجع"})  # the money went back after the statement was issued
    adj = cl.post(f"/api/lea/owners/{oid}/statements", json={"period": period}).json()
    assert adj["kind"] == "adjustment" and adj["net"] == -950  # collected −1000, commission −50, expenses unchanged
    rows_ = cl.get("/api/lea/statements", params={"owner_id": oid}).json()
    assert len(rows_) == 2
    final = cl.get(f"/api/lea/owners/{oid}/statement", params={"period": period}).json()
    assert round(sum(x["net"] for x in final["issued"]), 3) == final["net"] == -100  # the two statements add up to the month's true figure
    first = next(x for x in rows_ if x["kind"] == "statement")
    assert cl.get(f"/api/docs/owner_statement/{first['id']}.pdf").content[:4] == b"%PDF"
    assert cl.post(f"/api/lea/statements/{first['id']}/transferred", json={"ref": "TRF-77"}).status_code == 200
    assert cl.post(f"/api/lea/statements/{first['id']}/transferred", json={"ref": "TRF-78"}).status_code == 409
    login_as(cl, "finance")
    assert cl.post(f"/api/lea/owners/{oid}/statements", json={"period": period}).status_code == 403


# ----------------------------------------------------------------------------------------------------- property manager (scope leak test)
def _make_pm(cl, name):
    login_as(cl, "admin")
    r = cl.post("/api/users", json={"username": name, "name": "مدير عقاري للاختبار", "role": "property_manager"})
    assert r.status_code == 200, r.text
    _sql("UPDATE users SET must_change=0 WHERE username=?", (name,))
    return _sql("SELECT id FROM users WHERE username=?", (name,))[0][0]


def test_property_manager_sees_and_collects_only_inside_his_scope(cl):
    login_as(cl, "admin")
    a = cl.post("/api/lea/owners", json={"name": "مالك اختبار أ", "kind": "third_party"}).json()["id"]
    b = cl.post("/api/lea/owners", json={"name": "مالك اختبار ب", "kind": "third_party"}).json()["id"]
    la = _lease(cl, _unit(cl, a), start="2040-01-01", frequency=12)
    lb = _lease(cl, _unit(cl, b), start="2040-01-01", frequency=12)
    pm = _make_pm(cl, "pmt-one")
    nobody = _make_pm(cl, "pmt-none")
    assert cl.post("/api/lea/scopes", json={"user_id": pm}).status_code == 400  # an open scope is refused
    assert cl.post("/api/lea/scopes", json={"user_id": _sql("SELECT id FROM users WHERE username='admin'")[0][0], "owner_id": a}).status_code == 409  # only for property managers
    sc = cl.post("/api/lea/scopes", json={"user_id": pm, "owner_id": a})
    assert sc.status_code == 200
    login_as(cl, "pmt-one")
    assert cl.get("/api/lea/leases").status_code == 403  # no office access at all
    mine = cl.get("/api/pm/leases").json()
    assert isinstance(mine, list), mine
    assert [x["id"] for x in mine] == [la]
    assert cl.get(f"/api/pm/leases/{la}").status_code == 200
    assert cl.get(f"/api/pm/leases/{lb}").status_code == 404  # outside the scope: not even 403
    assert cl.post("/api/pm/payments", json={"lease_id": lb, "amount": 10, "method": "cash"}).status_code == 404
    assert cl.post("/api/pm/payments", json={"lease_id": la, "amount": 10, "method": "cheque"}).status_code == 422
    ok = cl.post("/api/pm/payments", json={"lease_id": la, "amount": 100, "method": "cash"})
    assert ok.status_code == 200 and ok.json()["receipt_no"].startswith("RC-")
    ov = cl.get("/api/pm/overview").json()
    assert ov["units"] == 1 and ov["collected_month"] == 100
    detail = cl.get(f"/api/pm/leases/{la}").json()
    assert "id_number" not in detail["tenant"] and "owner" not in detail
    login_as(cl, "pmt-none")
    assert cl.get("/api/pm/leases").json() == [] and cl.get("/api/pm/overview").json()["units"] == 0
    login_as(cl, "admin")
    assert nobody and cl.delete(f"/api/lea/scopes/{sc.json()['id']}").status_code == 200
    login_as(cl, "pmt-one")
    assert cl.get("/api/pm/leases").json() == []


# ----------------------------------------------------------------------------------------------------- import
def _xlsx(header, rows_):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows_:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _up(cl, kind, data, name, confirm=""):
    return cl.post(f"/api/lea/import/{kind}", files={"file": (name, data, "application/octet-stream")}, data={"confirm": confirm} if confirm else {})


def test_import_units_tenants_leases_payments_dry_run_commit_and_rerun(cl):
    login_as(cl, "admin")
    project = cl.get("/api/lea/units").json()["units"][0]["project"]
    # manual unit that the file will disagree with → a conflict, never an overwrite
    manual = _unit(cl)
    manual_code = _sql("SELECT code FROM units WHERE id=?", (manual,))[0][0]
    f = _xlsx(["المشروع", "رمز الوحدة", "النوع", "المساحة م²", "الاستخدام", "المواقف"],
              [[project, "T-OPS-IMP1", "مكتب", 80, "تجاري", 2], [project, "T-OPS-IMP2", "محل", 40, "commercial", None], [project, manual_code, "مستودع", 10, None, None]])
    before = _sql("SELECT COUNT(*) FROM units WHERE code LIKE 'T-OPS-%'")[0][0]
    dry = _up(cl, "units", f, "ops-test-units.xlsx").json()
    assert dry["mode"] == "dry-run" and dry["committed"] is False and dry["counts"]["insert"] == 2 and dry["counts"]["conflict"] == 1 and dry["errors"] == []
    assert _sql("SELECT COUNT(*) FROM units WHERE code LIKE 'T-OPS-%'")[0][0] == before  # a dry-run writes nothing
    assert _up(cl, "units", f, "ops-test-units.xlsx", "0" * 64).status_code == 409  # the confirmation must be the file's own hash
    done = _up(cl, "units", f, "ops-test-units.xlsx", dry["sha256"]).json()
    assert done["committed"] is True and _sql("SELECT COUNT(*) FROM units WHERE code LIKE 'T-OPS-%'")[0][0] == before + 2
    assert _sql("SELECT type FROM units WHERE id=?", (manual,))[0][0] == "مكتب"  # the manual row was not touched
    assert _up(cl, "units", f, "ops-test-units.xlsx", dry["sha256"]).status_code == 409  # same file twice
    f2 = _xlsx(["المشروع", "رمز الوحدة", "النوع", "المواقف"], [[project, "T-OPS-IMP1", "مكتب", 3]])  # a different file naming a unit the import made: fills only the empty
    again = _up(cl, "units", f2, "ops-test-units2.xlsx").json()
    assert again["counts"]["insert"] == 0 and again["counts"]["update"] + again["counts"]["skip"] == 1
    # tenants
    ft = _xlsx(["الاسم", "رقم الهوية", "الهاتف", "النوع (فرد/شركة)"], [["مستأجر استيراد الأول", "IM100200", "+968 90111222", "فرد"], ["مستأجر استيراد الثاني", "IM100201", "", "شركة"]])
    dt_ = _up(cl, "tenants", ft, "ops-test-tenants.xlsx").json()
    assert dt_["counts"]["insert"] == 2
    assert _up(cl, "tenants", ft, "ops-test-tenants.xlsx", dt_["sha256"]).json()["committed"] is True
    # leases: a signed historic one, a draft, and a bad row (unknown tenant)
    fl = _xlsx(["رمز الوحدة", "المستأجر (هوية/سجل/اسم)", "بداية العقد", "المدة بالأشهر", "الإيجار السنوي (السنة 1)", "عدد الدفعات سنويًا", "رقم العقد", "موقَّع (نعم/لا)"],
               [["T-OPS-IMP1", "IM100200", "2035-01-01", 12, 1200, 4, "OLD-2035-77", "نعم"], ["T-OPS-IMP2", "IM100201", "2035-01-01", 12, 600, 12, None, "لا"]])
    bad = _xlsx(["رمز الوحدة", "المستأجر (هوية/سجل/اسم)", "بداية العقد", "المدة بالأشهر", "الإيجار السنوي (السنة 1)"], [["T-OPS-IMP1", "NOBODY", "2036-01-01", 12, 900]])
    rb = _up(cl, "leases", bad, "ops-test-bad.xlsx").json()
    assert rb["errors"] and rb["errors"][0]["row"] == 2 and "المستأجر" in rb["errors"][0]["col"]
    assert _up(cl, "leases", bad, "ops-test-bad.xlsx", rb["sha256"]).json()["committed"] is False  # errors block the commit
    dl = _up(cl, "leases", fl, "ops-test-leases.xlsx").json()
    assert dl["counts"]["insert"] == 2 and dl["errors"] == []
    assert _up(cl, "leases", fl, "ops-test-leases.xlsx", dl["sha256"]).json()["committed"] is True
    leases = {x["number"]: x for x in cl.get("/api/lea/leases", params={"q": "T-OPS-IMP"}).json()}
    assert "OLD-2035-77" in leases and leases["OLD-2035-77"]["status"] == "active"
    imp = leases["OLD-2035-77"]
    assert len(_detail(cl, imp["id"])["dues"]) == 4 and any(x["status"] == "draft" for x in leases.values())
    # payments: history without receipts, never twice
    fp = _xlsx(["رقم العقد", "المبلغ", "تاريخ السداد", "تاريخ الاستحقاق", "الطريقة (نقد/شيك/تحويل/بطاقة)", "المرجع"],
               [["OLD-2035-77", 300, _past_start(5), "2035-01-01", "تحويل", "TRF-OLD-1"]])
    dp = _up(cl, "payments", fp, "ops-test-pay.xlsx").json()
    assert dp["counts"]["insert"] == 1 and dp["errors"] == []
    assert _up(cl, "payments", fp, "ops-test-pay.xlsx", dp["sha256"]).json()["committed"] is True
    fp2 = _xlsx(["رقم العقد", "المبلغ", "تاريخ السداد", "تاريخ الاستحقاق", "الطريقة (نقد/شيك/تحويل/بطاقة)", "المرجع", "ملاحظة"],
                [["OLD-2035-77", 300, _past_start(5), "2035-01-01", "تحويل", "TRF-OLD-1", "x"]])  # different file, same payment
    assert _up(cl, "payments", fp2, "ops-test-pay2.xlsx").json()["counts"] == {"insert": 0, "update": 0, "skip": 1, "conflict": 0}
    d = _detail(cl, imp["id"])
    assert d["dues"][0]["paid"] == 300 and d["payments"][0]["receipt_no"] is None and d["payments"][0]["method"] == "transfer"
    _no_mismatch()


def test_import_refuses_formulas_bad_formats_and_missing_columns(cl):
    login_as(cl, "admin")
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["الاسم", "رقم الهوية"])
    ws.append(["مستأجر استيراد صيغة", "=1+1"])
    ws["B2"] = "=SUM(1,2)"
    buf = io.BytesIO()
    wb.save(buf)
    assert _up(cl, "tenants", buf.getvalue(), "ops-test-f.xlsx").status_code == 400
    assert _up(cl, "tenants", b"x,y\n1,2\n", "ops-test-c.csv").status_code == 400  # required column missing
    assert _up(cl, "tenants", b"abc", "ops-test.txt").status_code == 400
    assert _up(cl, "tenants", b"\xff\xfe", "ops-test-b.csv").status_code == 400
    assert _up(cl, "nope", b"a", "ops-test-n.csv").status_code == 404
    csv = "الاسم,رقم الهوية\nمستأجر استيراد سي إس في,CS100\n".encode("utf-8-sig")
    assert _up(cl, "tenants", csv, "ops-test-ok.csv").json()["counts"]["insert"] == 1
    assert cl.get("/api/lea/import/leases/template.xlsx").content[:2] == b"PK"
    login_as(cl, "finance")
    assert _up(cl, "tenants", csv, "ops-test-ok.csv").status_code == 403  # importing is the manager's job


# ----------------------------------------------------------------------------------------------------- reports
def test_leasing_ready_reports_run_and_follow_the_ledger(cl):
    login_as(cl, "admin")
    lid = _lease(cl, start=_past_start(60), frequency=1)
    cl.post("/api/lea/payments", json={"lease_id": lid, "amount": 400, "method": "transfer"})
    keys = ("lease_monthly", "lease_occupancy", "lease_aging", "lease_expiring", "lease_cheques", "lease_owners")
    cat = {r["key"] for r in cl.get("/api/reports/catalog").json()["ready"]}
    assert set(keys) <= cat
    for k in keys:
        r = cl.get(f"/api/reports/ready/{k}")
        assert r.status_code == 200 and r.json()["kpis"] is not None and r.json()["tables"], (k, r.text[:200])
    m = cl.get("/api/reports/ready/lease_monthly").json()
    assert next(x for x in m["kpis"] if x["label"] == "المحصَّل خلال الشهر")["value"] >= 400
    assert any(t["rows"] for t in m["tables"])
    login_as(cl, "sales")
    assert cl.get("/api/reports/ready/lease_monthly").status_code in (403, 404)  # not for roles without the leasing permission
