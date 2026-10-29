"""0.24.0 — accounting, completion: supplier bills and payment orders (two approvers above the limit), Islamic-only bank accounts, statement import and reconciliation,
the VAT return, fixed assets and depreciation, budgets, Islamic financing with a fixed mark-up, the four financial statements and their exports, the year-end close,
and the ERPNext export. Every statement is checked against its own proof (balanced, reconciled)."""
import datetime as dt
import io
import os
import tempfile

os.environ.setdefault("MABANIQ_DB", os.path.join(tempfile.mkdtemp(), "accfull.db"))
os.environ["MABANIQ_FORCE_PW_CHANGE"] = "0"
os.environ["MABANIQ_TENANTS"] = "bayan:بيان عمان;almaqsura:المقصورة العقارية"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from backend import auth as A  # noqa: E402
from backend.app import app  # noqa: E402
from backend.db import TENANT, connect, init  # noqa: E402

CSRF = "csrf-accf-" + "f" * 30
TODAY = dt.date.today()
ACC_TABLES = ("fin_installments", "fin_contracts", "acc_budgets", "depreciation_runs", "fixed_assets", "vat_returns", "bank_recons", "bank_stmt_lines", "payment_orders", "ap_payments",
              "ap_bill_lines", "ap_bills", "bank_accounts", "acc_lines", "acc_journals", "acc_source_state", "acc_periods", "acc_cost_centers", "acc_accounts")
CODE = {"cash": "1110", "bank": "1121", "ar": "1210", "vat_in": "1240", "vendors": "2110", "vat_out": "2160", "vat_payable": "2165", "fin_liab": "2180", "deferred": "2185",
        "capital": "3100", "retained": "3300", "rent": "4110", "other_rev": "4190", "maint": "5110", "admin": "5230", "margin": "5260", "equip": "1520", "accum": "1590", "dep": "5240",
        "gain": "4195", "loss": "5270"}


def login_as(c, user, tenant="bayan"):
    tok = TENANT.set(tenant)
    try:
        c.cookies.set(A.COOKIE, A.issue_session(user))
    finally:
        TENANT.reset(tok)
    c.cookies.set(A.TENANT_COOKIE, tenant)
    c.cookies.set(A.CSRF_COOKIE, CSRF)
    c.headers["X-CSRF-Token"] = CSRF
    return c


def _sql(sql, args=(), maint=False):
    tok = TENANT.set("bayan")
    try:
        c = connect()
        try:
            pg = getattr(c, "dialect", "sqlite") == "postgres"
            if maint:
                c.maintenance()
            r = c.execute(sql, args)
            out = r.fetchall() if sql.lstrip().upper().startswith("SELECT") else None
            c.commit()
            if maint:
                c.end_maintenance()
                if not pg:
                    init(c)
            return out
        finally:
            c.close()
    finally:
        TENANT.reset(tok)


def _wipe():
    tok = TENANT.set("bayan")
    try:
        c = connect()
        pg = getattr(c, "dialect", "sqlite") == "postgres"
        sc = " AND org_id = current_setting('app.org_id', true)" if pg else ""
        c.maintenance()
        c.execute(f"DELETE FROM cheques WHERE direction='out'{sc}")
        for t in ACC_TABLES:
            c.execute(c.wipe_sql(t))
        c.execute(f"DELETE FROM settings WHERE k IN ('acc_start','acc_attach_over','ap_dual_over'){sc}")
        c.execute(f"DELETE FROM counters WHERE k LIKE 'jv:%' OR k LIKE 'apbill:%' OR k LIKE 'porder:%'{sc}")
        c.execute(f"DELETE FROM documents WHERE ref_type='journal'{sc}")
        c.execute(f"DELETE FROM vendors WHERE name LIKE 'مورّد محاسبة%'{sc}")
        c.execute(f"DELETE FROM sessions WHERE user_id IN (SELECT id FROM users WHERE username LIKE 'accf-%'){sc}")
        c.execute(f"DELETE FROM users WHERE username LIKE 'accf-%'{sc}")
        c.commit()
        c.end_maintenance()
        if not pg:
            init(c)
        c.close()
    finally:
        TENANT.reset(tok)


@pytest.fixture()
def cl():
    from backend import observability as O
    O.limiter.reset()
    with TestClient(app, base_url="https://testserver") as c:
        _wipe()
        for t in ("jadwa", "bayan", "almaqsura"):
            tok = TENANT.set(t)
            try:
                d = connect()
                d.execute("DELETE FROM login_attempts")
                d.commit()
                d.close()
            finally:
                TENANT.reset(tok)
        yield c


@pytest.fixture(scope="module", autouse=True)
def _cleanup():
    yield
    _wipe()


# ----------------------------------------------------------------------------------------------------- helpers
def _activate(cl, days_back=120, start=None):
    login_as(cl, "admin")
    s = start or (TODAY - dt.timedelta(days=days_back)).isoformat()
    assert cl.post("/api/acc/activate", json={"start": s, "confirm": True}).status_code == 200
    return s


def _vendor(cl, name="مورّد محاسبة", **kw):
    r = cl.post("/api/acc/vendors", json={"name": f"{name} {os.getpid()}", "tax_no": "OM1234567", **kw})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _bank_id(cl, kind="bank"):
    return next(b["id"] for b in cl.get("/api/acc/banks").json() if b["kind"] == kind)


def _bill(cl, vendor, lines, vat="out_of_scope", date=None, approve=True):
    r = cl.post("/api/acc/bills", json={"vendor_id": vendor, "bill_date": (date or TODAY - dt.timedelta(days=3)).isoformat(), "vat_class": vat, "vendor_ref": "INV-77",
                                       "lines": [{"account": a, "amount": m} for a, m in lines]})
    assert r.status_code == 200, r.text
    bid = r.json()["id"]
    if approve:
        a = cl.post(f"/api/acc/bills/{bid}/approve")
        assert a.status_code == 200, a.text
    return bid


def _je(cl, lines, date=None, memo="قيد اختبار"):
    """A manual journal, approved by the administrator (allowed to approve their own)."""
    r = cl.post("/api/acc/journals", json={"date": (date or TODAY).isoformat(), "memo": memo, "lines": [{"account": a, "debit": d, "credit": k} for a, d, k in lines]})
    assert r.status_code == 200, r.text
    ap = cl.post(f"/api/acc/journals/{r.json()['id']}/approve")
    assert ap.status_code == 200, ap.text
    return r.json()["id"]


def _tb(cl, **kw):
    r = cl.get("/api/acc/trial-balance", params={"frm": "2000-01-01", "to": TODAY.isoformat()} | kw)
    assert r.status_code == 200, r.text
    return {a["code"]: a for a in r.json()["accounts"]}


def _bal(cl, code):
    return _tb(cl).get(code, {"closing": 0})["closing"]


def _second_admin(cl, name="accf-admin2"):
    login_as(cl, "admin")
    r = cl.post("/api/users", json={"username": name, "name": "مدير ثانٍ للاختبار", "role": "admin"})
    assert r.status_code == 200, r.text
    _sql("UPDATE users SET must_change=0 WHERE username=?", (name,))
    return name


# ----------------------------------------------------------------------------------------------------- banks
def test_bank_accounts_are_islamic_only_and_the_event_accounts_get_their_records(cl):
    _activate(cl)
    banks = cl.get("/api/acc/banks").json()
    assert {b["kind"] for b in banks} >= {"cash", "bank", "charity"} and all(b["islamic"] == 1 for b in banks)
    bad = cl.post("/api/acc/banks", json={"name": "بنك تقليدي", "kind": "bank", "islamic": False})
    assert bad.status_code == 400 and "إسلامية" in bad.text
    ok = cl.post("/api/acc/banks", json={"name": "بنك إسلامي الثاني", "bank": "نافذة إسلامية", "kind": "bank", "islamic": True, "iban": "OM810180000001299123456"})
    assert ok.status_code == 200, ok.text
    new = next(b for b in cl.get("/api/acc/banks").json() if b["id"] == ok.json()["id"])
    assert new["iban"] and "3456" in new["iban"] and "OM81018" not in new["iban"]  # masked
    gl = next(a for a in cl.get("/api/acc/accounts").json() if a["id"] == ok.json()["account_id"])
    assert gl["cash_class"] == "cash" and gl["type"] == "asset" and gl["code"].startswith("11")
    assert cl.post("/api/acc/banks", json={"name": "ثالث", "islamic": True, "account_id": ok.json()["account_id"]}).status_code == 409  # that GL account already has a bank
    login_as(cl, "finance")
    assert cl.post("/api/acc/banks", json={"name": "x", "islamic": True}).status_code == 403
    with pytest.raises(Exception):  # noqa: B017 — the column itself can only ever hold 1 (CHECK)
        _sql("UPDATE bank_accounts SET islamic=0 WHERE id=(SELECT MIN(id) FROM bank_accounts)")


# ----------------------------------------------------------------------------------------------------- bills + payment orders
def test_a_supplier_bill_posts_expense_and_input_vat_against_the_payable_when_approved(cl):
    _activate(cl)
    v = _vendor(cl)
    login_as(cl, "finance")
    bid = _bill(cl, v, [(CODE["maint"], 400), (CODE["admin"], 100)], vat="standard", approve=False)
    b = cl.get(f"/api/acc/bills/{bid}").json()
    assert (b["net"], b["vat"], b["total"], b["status"]) == (500, 25, 525, "draft") and b["number"] is None
    assert cl.post(f"/api/acc/bills/{bid}/approve").status_code == 403  # the accountant cannot approve
    assert _bal(cl, CODE["vendors"]) == 0  # a draft posts nothing
    login_as(cl, "admin")
    r = cl.post(f"/api/acc/bills/{bid}/approve")
    assert r.status_code == 200 and r.json()["number"].startswith(f"AP-{TODAY.year}-"), r.text
    j = cl.get(f"/api/acc/journals/{r.json()['journal_id']}").json()
    got = {(ln["account_code"], ln["debit"], ln["credit"]) for ln in j["lines"]}
    assert got == {(CODE["maint"], 400, 0), (CODE["admin"], 100, 0), (CODE["vat_in"], 25, 0), (CODE["vendors"], 0, 525)}
    assert j["lines"][0]["vat_class"] == "standard" and any(ln["party_name"] for ln in j["lines"] if ln["account_code"] == CODE["vendors"])
    assert cl.post(f"/api/acc/bills/{bid}/approve").status_code == 409
    assert cl.put(f"/api/acc/bills/{bid}", json={"vendor_id": v, "bill_date": TODAY.isoformat(), "lines": [{"account": CODE["maint"], "amount": 1}]}).status_code == 409
    assert cl.post(f"/api/acc/bills/{bid}/void", json={"reason": "فاتورة مكررة"}).status_code == 200  # nothing paid yet ⟵ reversed
    assert _bal(cl, CODE["vendors"]) == 0 and _bal(cl, CODE["vat_in"]) == 0
    assert cl.get(f"/api/acc/bills/{bid}").json()["status"] == "void"
    assert cl.get("/api/acc/status").json()["ledger_diff"] == 0


def test_a_payment_order_pays_part_of_a_bill_and_a_reversal_gives_it_back(cl):
    _activate(cl)
    v = _vendor(cl)
    login_as(cl, "admin")
    bid = _bill(cl, v, [(CODE["maint"], 500)])
    bank = _bank_id(cl)
    login_as(cl, "finance")
    o = cl.post("/api/acc/orders", json={"bill_id": bid, "amount": 200, "method": "transfer", "bank_account_id": bank, "reference": "TRF-1"})
    assert o.status_code == 200, o.text
    assert cl.post(f"/api/acc/orders/{o.json()['id']}/execute", json={}).status_code == 409  # not approved yet
    assert cl.post("/api/acc/orders", json={"bill_id": bid, "amount": 400, "method": "transfer", "bank_account_id": bank}).status_code == 409  # 200 pending + 400 > 500
    login_as(cl, "admin")
    assert cl.post(f"/api/acc/orders/{o.json()['id']}/approve").json()["status"] == "approved"  # below the limit: one approver
    login_as(cl, "finance")
    ex = cl.post(f"/api/acc/orders/{o.json()['id']}/execute", json={})
    assert ex.status_code == 200, ex.text
    b = cl.get(f"/api/acc/bills/{bid}").json()
    assert (b["paid"], b["status"], b["outstanding"]) == (200, "partly_paid", 300)
    assert _bal(cl, CODE["vendors"]) == -300 and _bal(cl, CODE["bank"]) == -200
    assert cl.post(f"/api/acc/orders/{o.json()['id']}/execute", json={}).status_code == 409  # once
    login_as(cl, "admin")
    pid = cl.get(f"/api/acc/bills/{bid}").json()["payments"][0]["id"]
    assert cl.post(f"/api/acc/payments/{pid}/reverse", json={"reason": "تحويل مرفوض من البنك"}).status_code == 200
    b = cl.get(f"/api/acc/bills/{bid}").json()
    assert (b["paid"], b["status"]) == (0, "approved") and len(b["payments"]) == 2 and b["payments"][1]["amount"] == -200
    assert _bal(cl, CODE["vendors"]) == -500 and _bal(cl, CODE["bank"]) == 0
    assert cl.post(f"/api/acc/payments/{pid}/reverse", json={"reason": "مرة ثانية"}).status_code == 409
    if os.environ.get("MABANIQ_DATABASE_URL"):  # on PostgreSQL a supplier payment is a ledger row: the application role has no UPDATE on it
        with pytest.raises(Exception):  # noqa: B017
            _sql("UPDATE ap_payments SET amount=1 WHERE id=?", (pid,))


def test_above_the_limit_two_different_approvers_are_needed_and_never_the_creator(cl):
    _activate(cl)
    v = _vendor(cl)
    second = _second_admin(cl)
    login_as(cl, "admin")
    _sql("INSERT OR REPLACE INTO settings VALUES(?,?)", ("ap_dual_over", "100"))
    bid = _bill(cl, v, [(CODE["maint"], 900)])
    bank = _bank_id(cl)
    login_as(cl, "finance")
    oid = cl.post("/api/acc/orders", json={"bill_id": bid, "amount": 300, "method": "cheque", "bank_account_id": bank, "cheque_number": "C-9001", "cheque_due": TODAY.isoformat()}).json()["id"]
    assert cl.get("/api/acc/orders").json()[0]["needs_second"] is True
    login_as(cl, "admin")
    assert cl.post(f"/api/acc/orders/{oid}/approve").json()["status"] == "pending"  # first approval: still waiting for the second
    assert cl.post(f"/api/acc/orders/{oid}/approve").status_code == 403  # the same person cannot approve twice
    assert cl.post(f"/api/acc/orders/{oid}/execute", json={}).status_code == 409
    login_as(cl, second)
    assert cl.post(f"/api/acc/orders/{oid}/approve").json()["status"] == "approved"
    login_as(cl, "admin")
    r = cl.post(f"/api/acc/orders/{oid}/execute", json={})
    assert r.status_code == 200 and r.json()["cheque_id"], r.text
    out = cl.get("/api/acc/cheques-out").json()
    q = next(x for x in out if x["id"] == r.json()["cheque_id"])
    assert q["number"] == "C-9001" and q["status"] == "issued" and q["amount"] == 300
    assert cl.post(f"/api/acc/cheques-out/{q['id']}/cleared", json={}).status_code == 200
    assert cl.post(f"/api/acc/cheques-out/{q['id']}/cleared", json={}).status_code == 409
    # an order made by an administrator is approved by two OTHER people: the creator is never an approver
    login_as(cl, "admin")
    o2 = cl.post("/api/acc/orders", json={"bill_id": bid, "amount": 200, "method": "transfer", "bank_account_id": bank}).json()["id"]
    assert cl.post(f"/api/acc/orders/{o2}/approve").status_code == 403
    login_as(cl, second)
    assert cl.post(f"/api/acc/orders/{o2}/reject", json={"note": "مبلغ غير مبرَّر"}).status_code == 200
    assert cl.post(f"/api/acc/orders/{o2}/approve").status_code == 409


def test_aging_of_what_is_owed_to_suppliers(cl):
    _activate(cl)
    v = _vendor(cl)
    _bill(cl, v, [(CODE["maint"], 100)], date=TODAY - dt.timedelta(days=70))
    _bill(cl, v, [(CODE["admin"], 50)], date=TODAY - dt.timedelta(days=5))
    ag = cl.get("/api/acc/ap-aging").json()
    assert ag["total"] == 150 and ag["buckets"]["d61_90"] == 100 and ag["buckets"]["d1_30"] == 50


# ----------------------------------------------------------------------------------------------------- bank statement + reconciliation
def _csv(rows_):
    return ("التاريخ,البيان,المرجع,المبلغ,الرصيد\n" + "\n".join(",".join(str(x) for x in r) for r in rows_)).encode("utf-8-sig")


def test_a_statement_is_imported_once_matched_and_the_reconciliation_closes_only_when_it_balances(cl):
    _activate(cl)
    bank = _bank_id(cl)
    d = lambda n: (TODAY - dt.timedelta(days=n)).isoformat()  # noqa: E731
    _je(cl, [(CODE["bank"], 500, 0), (CODE["capital"], 0, 500)], TODAY - dt.timedelta(days=20), "إيداع رأس مال")
    _je(cl, [(CODE["admin"], 30, 0), (CODE["bank"], 0, 30)], TODAY - dt.timedelta(days=10), "رسوم")
    _je(cl, [(CODE["bank"], 70, 0), (CODE["other_rev"], 0, 70)], TODAY - dt.timedelta(days=2), "إيداع لم يصل البنك بعد")
    file = _csv([(d(20), "إيداع", "R1", 500, 500), (d(9), "رسوم بنكية", "R2", -30, 470), (d(5), "رسوم دفتر لم تُسجَّل", "R3", -5, 465)])
    up = cl.post(f"/api/acc/banks/{bank}/import", files={"file": ("stmt.csv", file, "text/csv")})
    assert up.status_code == 200 and up.json()["new"] == 3, up.text
    again = cl.post(f"/api/acc/banks/{bank}/import", files={"file": ("stmt.csv", file, "text/csv")}).json()
    assert again["new"] == 0 and again["duplicate"] == 3  # importing the same file twice adds nothing
    m = cl.post(f"/api/acc/banks/{bank}/auto-match").json()
    assert m["matched"] == 2 and m["ambiguous"] == 0
    lines = cl.get(f"/api/acc/banks/{bank}/lines").json()
    left = [x for x in lines if x["status"] == "unmatched"]
    assert len(left) == 1 and left[0]["amount"] == -5
    end = TODAY.isoformat()
    st = cl.get(f"/api/acc/banks/{bank}/recon", params={"end": end, "statement_balance": 465}).json()
    assert st["book_balance"] == 540 and st["in_transit"] == 70 and st["not_in_books"] == -5 and st["difference"] == 0 and st["balanced"]
    assert cl.post(f"/api/acc/banks/{bank}/recon/close", json={"end": end, "statement_balance": 470}).status_code == 409  # a wrong statement balance does not close
    assert cl.post(f"/api/acc/banks/{bank}/recon/close", json={"end": end, "statement_balance": 465}).status_code == 200
    assert cl.post(f"/api/acc/banks/{bank}/recon/close", json={"end": end, "statement_balance": 465}).status_code == 409
    assert cl.post(f"/api/acc/banks/{bank}/lines/{[x for x in lines if x['status'] == 'matched'][0]['id']}/unmatch").status_code == 409  # frozen by the closed reconciliation
    # a manual match needs the same amount
    sid = left[0]["id"]
    cand = cl.get(f"/api/acc/banks/{bank}/suggest/{sid}").json()
    assert cand and cl.post(f"/api/acc/banks/{bank}/lines/{sid}/match", json={"acc_line_id": cand[0]["id"]}).status_code == 409  # 70 ≠ −5


def test_two_equal_candidates_are_left_for_a_person_not_guessed(cl):
    _activate(cl)
    bank = _bank_id(cl)
    day = TODAY - dt.timedelta(days=4)
    for _ in range(2):
        _je(cl, [(CODE["bank"], 40, 0), (CODE["other_rev"], 0, 40)], day)
    up = cl.post(f"/api/acc/banks/{bank}/import", files={"file": ("s.csv", _csv([(day.isoformat(), "إيداع", "X", 40, 40)]), "text/csv")})
    assert up.status_code == 200
    m = cl.post(f"/api/acc/banks/{bank}/auto-match").json()
    assert m["matched"] == 0 and m["ambiguous"] == 1


# ----------------------------------------------------------------------------------------------------- VAT
def _prev_quarter():
    q = (TODAY.month - 1) // 3  # 0..3 → the previous quarter is q (1-based) of this year, or Q4 of last year
    year, qn = (TODAY.year, q) if q >= 1 else (TODAY.year - 1, 4)
    first = dt.date(year, 3 * (qn - 1) + 1, 1)
    return year, qn, first + dt.timedelta(days=12)


def test_the_vat_return_is_computed_from_the_ledger_filed_once_and_settles_the_vat_accounts(cl):
    start = (_prev_quarter()[2] - dt.timedelta(days=30)).isoformat()
    _activate(cl, start=start)
    year, q, day = _prev_quarter()
    v = _vendor(cl)
    _bill(cl, v, [(CODE["admin"], 500)], vat="standard", date=day)  # input VAT 25
    _je(cl, [(CODE["bank"], 420, 0), (CODE["other_rev"], 0, 400), (CODE["vat_out"], 0, 20)], day, "مبيعات خاضعة")  # output VAT 20
    r = cl.get(f"/api/acc/vat/{year}/{q}").json()
    assert (r["output_vat"], r["input_vat"], r["net"]) == (20, 25, -5) and r["boxes"]["purchases_standard"] == 500 and not r["filed"]
    assert cl.post(f"/api/acc/vat/{TODAY.year}/{(TODAY.month - 1) // 3 + 1}/file", json={"ref": "ACK-0"}).status_code == 409  # the quarter has not ended
    f = cl.post(f"/api/acc/vat/{year}/{q}/file", json={"ref": "ACK-2026-Q"})
    assert f.status_code == 200 and f.json()["net"] == -5, f.text
    assert _bal(cl, CODE["vat_out"]) == 0 and _bal(cl, CODE["vat_in"]) == 0 and _bal(cl, "1245") == 5  # refundable: input exceeded output
    again = cl.post(f"/api/acc/vat/{year}/{q}/file", json={"ref": "ACK-2"})
    assert again.status_code == 409
    assert cl.get(f"/api/acc/vat/{year}/{q}").json()["output_vat"] == 20  # the settlement journal is not counted in the quarter it settles
    assert len(cl.get("/api/acc/vat").json()) == 1


# ----------------------------------------------------------------------------------------------------- fixed assets, budgets
def test_straight_line_depreciation_is_idempotent_stops_at_the_salvage_and_disposal_books_the_gain(cl):
    _activate(cl, days_back=200)
    acq = TODAY - dt.timedelta(days=100)
    _je(cl, [(CODE["equip"], 900, 0), (CODE["capital"], 0, 900)], acq - dt.timedelta(days=1), "أصل ثابت عيني")
    r = cl.post("/api/acc/assets", json={"code": f"AST-{os.getpid()}", "name": "مكيفات", "cost": 900, "acquired": acq.isoformat(), "life_months": 9, "salvage": 90})
    assert r.status_code == 200, r.text
    aid = r.json()["id"]
    assert cl.post("/api/acc/assets", json={"code": f"AST-{os.getpid()}", "name": "مكرر", "cost": 1, "acquired": acq.isoformat(), "life_months": 1}).status_code == 409
    assert cl.post("/api/acc/assets", json={"code": "BAD-1", "name": "سالب", "cost": 10, "acquired": acq.isoformat(), "life_months": 1, "salvage": 10}).status_code == 400
    p0 = f"{acq.year}-{acq.month:02d}"
    assert cl.post("/api/acc/depreciation/run", json={"period": p0}).json()["assets"] == 0  # not in the month of acquisition
    nxt = (acq.replace(day=1) + dt.timedelta(days=32)).replace(day=1)
    p1 = f"{nxt.year}-{nxt.month:02d}"
    run = cl.post("/api/acc/depreciation/run", json={"period": p1}).json()
    assert run["assets"] == 1 and run["total"] == 90  # (900 − 90) / 9
    assert cl.post("/api/acc/depreciation/run", json={"period": p1}).json()["assets"] == 0  # idempotent
    assert _bal(cl, CODE["accum"]) == -90 and _bal(cl, CODE["dep"]) == 90
    asset = next(a for a in cl.get("/api/acc/assets").json() if a["id"] == aid)
    assert asset["accumulated"] == 90 and asset["book_value"] == 810 and asset["monthly"] == 90
    assert cl.post("/api/acc/depreciation/run", json={"period": f"{TODAY.year + 1}-01"}).status_code == 400
    bank = _bank_id(cl)
    d = cl.post(f"/api/acc/assets/{aid}/dispose", json={"on": TODAY.isoformat(), "proceeds": 1000, "bank_account_id": bank})
    assert d.status_code == 200 and d.json()["book_value"] == 810 and d.json()["gain_loss"] == 190, d.text
    assert _bal(cl, CODE["equip"]) == 0 and _bal(cl, CODE["accum"]) == 0 and _bal(cl, CODE["gain"]) == -190
    assert cl.post(f"/api/acc/assets/{aid}/dispose", json={"on": TODAY.isoformat()}).status_code == 409


def test_budget_against_actual_reports_variance_per_account(cl):
    _activate(cl, days_back=100)
    v = _vendor(cl)
    month = TODAY.month
    _bill(cl, v, [(CODE["maint"], 300)], date=TODAY)
    r = cl.post("/api/acc/budgets", json={"year": TODAY.year, "lines": [{"account": CODE["maint"], "month": month, "amount": 250}, {"account": CODE["admin"], "month": month, "amount": 100}]})
    assert r.status_code == 200 and r.json()["cells"] == 2
    cl.post("/api/acc/budgets", json={"year": TODAY.year, "lines": [{"account": CODE["admin"], "month": month, "amount": 0}]})  # a zero removes the cell
    rep = cl.get("/api/acc/budget-report", params={"year": TODAY.year, "to_month": month}).json()
    row = next(x for x in rep["rows"] if x["code"] == CODE["maint"])
    assert (row["budget"], row["actual"], row["variance"], row["pct"]) == (250, 300, -50, 120.0)  # over budget ⟵ negative variance
    assert not [x for x in rep["rows"] if x["code"] == CODE["admin"]]


# ----------------------------------------------------------------------------------------------------- Islamic financing
def test_financing_has_a_fixed_known_markup_defers_it_and_recognises_it_with_each_instalment(cl):
    _activate(cl)
    bank = _bank_id(cl)
    base = {"kind": "murabaha", "counterparty": "بنك إسلامي", "principal": 1000, "markup_total": 120, "start_date": (TODAY - dt.timedelta(days=60)).isoformat(), "bank_account_id": bank}
    inst = [{"due": (TODAY + dt.timedelta(days=30 * i)).isoformat(), "amount": 280} for i in range(4)]
    assert cl.post("/api/acc/finance", json=base | {"installments": inst[:3]}).status_code == 409  # 3 × 280 ≠ 1000 + 120
    assert cl.post("/api/acc/finance", json=base | {"installments": inst, "rate": 0.05}).status_code == 422  # there is no rate field: a rate is refused
    assert cl.post("/api/acc/finance", json=base | {"installments": [dict(x, interest=3) for x in inst]}).status_code == 422
    r = cl.post("/api/acc/finance", json=base | {"installments": inst})
    assert r.status_code == 200 and r.json()["total"] == 1120, r.text
    assert _bal(cl, CODE["bank"]) == 1000 and _bal(cl, CODE["fin_liab"]) == -1120 and _bal(cl, CODE["deferred"]) == 120 and _bal(cl, CODE["margin"]) == 0
    fid = r.json()["id"]
    p = cl.post(f"/api/acc/finance/{fid}/installments/1/pay", json={})
    assert p.status_code == 200, p.text
    assert _bal(cl, CODE["fin_liab"]) == -840 and _bal(cl, CODE["deferred"]) == 90 and _bal(cl, CODE["margin"]) == 30 and _bal(cl, CODE["bank"]) == 720
    assert cl.post(f"/api/acc/finance/{fid}/installments/1/pay", json={}).status_code == 409
    for n in (2, 3, 4):
        assert cl.post(f"/api/acc/finance/{fid}/installments/{n}/pay", json={}).status_code == 200
    c = next(x for x in cl.get("/api/acc/finance").json() if x["id"] == fid)
    assert c["status"] == "closed" and c["remaining"] == 0 and round(sum(i["margin"] for i in c["installments"]), 3) == 120
    assert _bal(cl, CODE["fin_liab"]) == 0 and _bal(cl, CODE["deferred"]) == 0 and _bal(cl, CODE["margin"]) == 120


# ----------------------------------------------------------------------------------------------------- statements + exports
def _activity(cl):
    """Capital, a purchased asset, a sale with VAT, a bill, a depreciation run and a financing contract: every statement has something to say."""
    v = _vendor(cl)
    _je(cl, [(CODE["bank"], 900, 0), (CODE["capital"], 0, 900)], TODAY - dt.timedelta(days=80), "رأس المال")
    _je(cl, [(CODE["equip"], 600, 0), (CODE["bank"], 0, 600)], TODAY - dt.timedelta(days=70), "شراء أصل")
    _je(cl, [(CODE["bank"], 500, 0), (CODE["other_rev"], 0, 500)], TODAY - dt.timedelta(days=40), "إيراد")
    _bill(cl, v, [(CODE["maint"], 200)], date=TODAY - dt.timedelta(days=30))
    cl.post("/api/acc/assets", json={"code": f"ST-{os.getpid()}", "name": "معدات", "cost": 600, "acquired": (TODAY - dt.timedelta(days=70)).isoformat(), "life_months": 12})
    nxt = (TODAY.replace(day=1) - dt.timedelta(days=1))
    cl.post("/api/acc/depreciation/run", json={"period": f"{nxt.year}-{nxt.month:02d}"})
    bank = _bank_id(cl)
    cl.post("/api/acc/finance", json={"kind": "ijara", "counterparty": "مصرف", "principal": 300, "markup_total": 30, "start_date": (TODAY - dt.timedelta(days=20)).isoformat(),
                                     "bank_account_id": bank, "installments": [{"due": (TODAY + dt.timedelta(days=30)).isoformat(), "amount": 330}]})


def test_the_four_statements_carry_their_own_proof(cl):
    _activate(cl)
    _activity(cl)
    frm = (TODAY - dt.timedelta(days=100)).isoformat()
    bs = cl.get("/api/acc/statements/balance_sheet", params={"to": TODAY.isoformat()}).json()
    assert bs["balanced"] is True and bs["assets"] == round(bs["liabilities"] + bs["equity"], 3) and bs["earnings"] == round(500 - 200 - _dep_total(cl), 3)
    inc = cl.get("/api/acc/statements/income_statement", params={"frm": frm, "to": TODAY.isoformat()}).json()
    assert inc["revenue"] == 500 and inc["expenses"] == 200 + _dep_total(cl) and inc["net"] == round(500 - 200 - _dep_total(cl), 3)
    assert inc["notes"] and "التزام لا إيراد" in inc["notes"][0]["text"]
    cf = cl.get("/api/acc/statements/cash_flow", params={"frm": frm, "to": TODAY.isoformat()}).json()
    assert cf["reconciled"] is True and cf["difference"] == 0
    assert cf["investing"] == -600 and cf["financing"] == 900 + 300 + 30 - 30 and cf["net_profit"] == inc["net"]  # asset bought; capital + financing received (margin deferred nets out)
    assert cf["cash_close"] - cf["cash_open"] == cf["change"]
    eq = cl.get("/api/acc/statements/equity_changes", params={"frm": frm, "to": TODAY.isoformat()}).json()
    assert eq["reconciled"] is True and eq["closing"] == bs["equity"] and eq["net_profit"] == inc["net"] and eq["other"] == 900
    assert cl.get("/api/acc/statements/nonsense").status_code == 404


def _dep_total(cl):
    return round(sum(a["accumulated"] for a in cl.get("/api/acc/assets").json()), 3)


def test_the_statements_export_to_csv_xlsx_and_pdf(cl):
    _activate(cl)
    _activity(cl)
    q = {"to": TODAY.isoformat()}
    csv_ = cl.get("/api/acc/statements/balance_sheet", params=q | {"fmt": "csv"})
    assert csv_.status_code == 200 and csv_.content[:3] == b"\xef\xbb\xbf" and "قائمة المركز المالي" in csv_.text
    xl = cl.get("/api/acc/statements/income_statement", params=q | {"fmt": "xlsx"})
    assert xl.status_code == 200 and xl.content[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(xl.content)).active
    assert ws.sheet_view.rightToLeft and any(c.value == "صافي الربح (الخسارة)" for c in ws["A"])
    for kind in ("balance_sheet", "income_statement", "cash_flow", "equity_changes"):
        pdf = cl.get(f"/api/acc/statements/{kind}", params=q | {"fmt": "pdf"})
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF", kind
    login_as(cl, "sales")
    assert cl.get("/api/acc/statements/balance_sheet").status_code == 403


# ----------------------------------------------------------------------------------------------------- year-end close + ERPNext
def test_the_year_close_carries_profit_to_retained_earnings_and_a_closed_year_still_shows_its_profit(cl):
    y = TODAY.year - 1
    _activate(cl, start=f"{y}-01-01")
    _je(cl, [(CODE["bank"], 800, 0), (CODE["other_rev"], 0, 800)], dt.date(y, 5, 10), "إيراد السنة")
    _je(cl, [(CODE["admin"], 300, 0), (CODE["bank"], 0, 300)], dt.date(y, 8, 20), "مصروف السنة")
    assert cl.post("/api/acc/year-close", json={"year": TODAY.year}).status_code == 409  # a year still running
    r = cl.post("/api/acc/year-close", json={"year": y, "ignore_pending": True})
    assert r.status_code == 200 and r.json()["net_profit"] == 500, r.text
    assert cl.post("/api/acc/year-close", json={"year": y, "ignore_pending": True}).status_code == 409
    assert _bal(cl, CODE["retained"]) == -500 and _bal(cl, CODE["other_rev"]) == 0 and _bal(cl, CODE["admin"]) == 0
    inc = cl.get("/api/acc/statements/income_statement", params={"frm": f"{y}-01-01", "to": f"{y}-12-31"}).json()
    assert inc["net"] == 500  # the closing journal is invisible to the income statement
    bs = cl.get("/api/acc/statements/balance_sheet", params={"to": TODAY.isoformat()}).json()
    assert bs["balanced"] and bs["equity"] == 500 and bs["earnings"] == 0
    assert all(p["status"] == "closed" for p in cl.get("/api/acc/periods").json() if p["year"] == y)
    late = cl.post("/api/acc/journals", json={"date": dt.date(y, 12, 30).isoformat(), "memo": "قيد متأخر", "lines": [{"account": CODE["cash"], "debit": 5}, {"account": CODE["capital"], "credit": 5}]})
    jid = late.json()["id"]
    assert cl.post(f"/api/acc/journals/{jid}/approve").status_code == 409  # the year is closed
    assert cl.post(f"/api/acc/year-close/{y}/reopen", json={"reason": "x"}).status_code == 422
    assert cl.post(f"/api/acc/year-close/{y}/reopen", json={"reason": "تصحيح فاتورة سنة سابقة"}).status_code == 200
    assert _bal(cl, CODE["retained"]) == 0 and _bal(cl, CODE["other_rev"]) == -800
    assert cl.get("/api/acc/status").json()["ledger_diff"] == 0


def test_the_erpnext_export_lists_the_chart_and_the_journals_in_journal_entry_shape(cl):
    _activate(cl)
    _je(cl, [(CODE["cash"], 60, 0), (CODE["capital"], 0, 60)], TODAY - dt.timedelta(days=3), "تصدير")
    x = cl.get("/api/acc/export/erpnext", params={"frm": (TODAY - dt.timedelta(days=10)).isoformat()}).json()
    assert x["company_currency"] == "OMR" and {a["root_type"] for a in x["chart_of_accounts"]} == {"Asset", "Liability", "Equity", "Income", "Expense"}
    je = next(e for e in x["journal_entries"] if e["user_remark"] == "تصدير")
    assert je["voucher_type"] == "Journal Entry" and je["total_debit"] == je["total_credit"] == 60 and len(je["accounts"]) == 2
    assert {a["account_number"] for a in je["accounts"]} == {CODE["cash"], CODE["capital"]}
