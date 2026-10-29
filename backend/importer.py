"""Mabaniq 0.22.0 — Excel / CSV import of the leasing register (design §3.4).

The path is always: upload → dry-run (reads and checks, writes NOTHING) → a report of errors / warnings / conflicts → the same file again with its SHA-256 as the
confirmation → one transaction. Kinds: units, tenants, leases, payments, cheques — load them in that order (a lease needs its unit and tenant to exist).

Rules kept on purpose:
* It never wipes and never re-seeds. A new row is inserted; a row that already exists is touched only if an earlier import created it (and then only its empty
  fields are filled). An existing hand-made row that differs from the file is a *conflict*: reported, not applied.
* Re-running a file changes nothing: the natural keys (unit code, tenant id / commercial register / name+phone, lease number, payment fingerprint, cheque number+bank)
  find what is already there; a file already imported (same SHA-256) is refused outright.
* The file is data, not code: formulas are rejected (values only), at most 5 MB and 5,000 rows, a zip bomb is refused before it is opened.
* Imported payments are history: they carry no receipt number (the office issued its own at the time) and are stamped with a fingerprint so they cannot double.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import zipfile

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response

from . import pii
from .common import act_as, audit, db, now_s, rows, today
from .leasing import EPS, TENANT_SQL_COLS, USES, LeaseNew, TenantIn, _overlap, _tenant_cols, addm, apply_payment, create_lease, next_no, r3, sign_lease

router = APIRouter()

MAX_BYTES, MAX_ROWS, MAX_UNZIPPED = 5_000_000, 5000, 40_000_000
# kind → [(key, Arabic header, required)]
SPECS: dict[str, list[tuple[str, str, bool]]] = {
    "units": [("project", "المشروع", True), ("code", "رمز الوحدة", True), ("type", "النوع", True), ("area", "المساحة م²", False), ("use", "الاستخدام", False),
              ("owner", "المالك", False), ("parking_bays", "المواقف", False)],
    "tenants": [("name", "الاسم", True), ("kind", "النوع (فرد/شركة)", False), ("phone", "الهاتف", False), ("id_number", "رقم الهوية", False), ("cr_number", "السجل التجاري", False),
                ("email", "البريد", False), ("address", "العنوان", False), ("nationality", "الجنسية", False)],
    "leases": [("unit_code", "رمز الوحدة", True), ("tenant", "المستأجر (هوية/سجل/اسم)", True), ("start", "بداية العقد", True), ("months", "المدة بالأشهر", True),
               ("annual_rent", "الإيجار السنوي (السنة 1)", False), ("rent_y2", "إيجار السنة 2", False), ("rent_y3", "إيجار السنة 3", False), ("frequency", "عدد الدفعات سنويًا", False),
               ("prep_free_days", "أيام التجهيز المجانية", False), ("deposit", "التأمين", False), ("number", "رقم العقد", False), ("municipality_ref", "مرجع البلدية", False),
               ("signed", "موقَّع (نعم/لا)", False)],
    "payments": [("lease_number", "رقم العقد", True), ("amount", "المبلغ", True), ("paid_on", "تاريخ السداد", True), ("due_date", "تاريخ الاستحقاق", False),
                 ("method", "الطريقة (نقد/شيك/تحويل/بطاقة)", False), ("reference", "المرجع", False)],
    "cheques": [("lease_number", "رقم العقد", True), ("number", "رقم الشيك", True), ("bank", "البنك", True), ("amount", "المبلغ", True), ("due_date", "تاريخ الاستحقاق", True),
                ("received_on", "تاريخ الاستلام", False), ("drawer", "الساحب", False), ("status", "الحالة (محفوظ/مودَع)", False)],
}
KIND_LABEL = {"units": "الوحدات", "tenants": "المستأجرون", "leases": "العقود", "payments": "المدفوعات", "cheques": "الشيكات"}
METHOD_WORDS = {"نقد": "cash", "نقدا": "cash", "نقدًا": "cash", "cash": "cash", "شيك": "cheque", "cheque": "cheque", "تحويل": "transfer", "transfer": "transfer",
                "بطاقة": "card", "card": "card", "أخرى": "other", "other": "other"}
YES = {"نعم", "yes", "y", "1", "true", "صحيح"}


class Bad(Exception):
    """A row-level problem: (column key, reason)."""

    def __init__(self, col: str, reason: str):
        self.col, self.reason = col, reason


# ----------------------------------------------------------------------------------------------------- reading the file
def _cell(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v.strip() if isinstance(v, str) else v


def read_table(filename: str, data: bytes, kind: str) -> list[tuple[int, dict]]:
    """→ [(row number in the sheet, {key: value})]; raises HTTPException for a file that cannot be read at all."""
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "الملف أكبر من 5 ميغابايت")
    spec = SPECS[kind]
    names = {}
    for key, label, _req in spec:
        names[label.strip().lower()] = key
        names[key.lower()] = key
    if filename.lower().endswith(".csv"):
        try:
            grid = [list(r) for r in csv.reader(io.StringIO(data.decode("utf-8-sig")))]
        except UnicodeDecodeError:
            raise HTTPException(400, "ملف CSV يجب أن يكون بترميز UTF-8") from None
    elif filename.lower().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                    raise HTTPException(413, "محتوى الملف بعد الفك أكبر من المسموح")
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
        except HTTPException:
            raise
        except Exception:  # noqa: BLE001 — a corrupt workbook must not leak a traceback
            raise HTTPException(400, "تعذّرت قراءة ملف Excel") from None
        ws = wb.worksheets[0]
        grid = []
        for r in ws.iter_rows():
            if len(grid) > MAX_ROWS + 1:
                break
            line = []
            for cell in r:
                if cell.data_type == "f":
                    raise HTTPException(400, f"الصيغ والمعادلات ممنوعة (الخلية {cell.coordinate}) — الصق القيم فقط")
                line.append(_cell(cell.value))
            grid.append(line)
    else:
        raise HTTPException(400, "الصيغ المقبولة: xlsx أو csv")
    grid = [[_cell(v) for v in r] for r in grid]
    while grid and not any(v not in (None, "") for v in grid[0]):
        grid.pop(0)
    if not grid:
        raise HTTPException(400, "الملف فارغ")
    head = [names.get(str(h or "").strip().lower()) for h in grid[0]]
    missing = [label for key, label, req in spec if req and key not in head]
    if missing:
        raise HTTPException(400, "أعمدة مطلوبة غير موجودة: " + "، ".join(missing))
    out = []
    for i, r in enumerate(grid[1:], start=2):
        if not any(v not in (None, "") for v in r):
            continue
        out.append((i, {k: (r[j] if j < len(r) and r[j] != "" else None) for j, k in enumerate(head) if k}))
    if len(out) > MAX_ROWS:
        raise HTTPException(413, f"أكثر من {MAX_ROWS} صف")
    return out


# ----------------------------------------------------------------------------------------------------- value helpers
def _s(v, col, req=False, maxlen=120):
    v = None if v is None else str(v).strip()
    if v in (None, ""):
        if req:
            raise Bad(col, "مطلوب")
        return None
    if len(v) > maxlen:
        raise Bad(col, f"أطول من {maxlen} حرفًا")
    return v


def _n(v, col, req=False, lo=0.0, hi=10_000_000.0):
    if v in (None, ""):
        if req:
            raise Bad(col, "مطلوب")
        return None
    try:
        x = float(str(v).replace(",", ""))
    except ValueError:
        raise Bad(col, "ليس رقمًا") from None
    if not lo <= x <= hi:
        raise Bad(col, f"خارج المدى ({lo:g}–{hi:g})")
    return x


def _d(v, col, req=False):
    if v in (None, ""):
        if req:
            raise Bad(col, "مطلوب")
        return None
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise Bad(col, "تاريخ غير صالح (استعمل 2026-01-31)")


# ----------------------------------------------------------------------------------------------------- planning (read-only)
class Ctx:
    """Everything a plan needs, read once: lookups of the existing register. Never writes."""

    def __init__(self, c, kind: str):
        self.c, self.kind = c, kind
        self.projects = {r["name"].strip(): r["id"] for r in c.execute("SELECT id, name FROM projects").fetchall()}
        self.owners = {r["name"].strip(): r["id"] for r in c.execute("SELECT id, name FROM owners WHERE COALESCE(active,1)=1").fetchall()}
        own = c.execute("SELECT id FROM owners WHERE kind='own' ORDER BY id LIMIT 1").fetchone()
        self.own = own["id"] if own else None
        self.seen: set = set()


def _plan_units(x: Ctx, n: int, r: dict) -> dict:
    proj = _s(r.get("project"), "project", True)
    pid = x.projects.get(proj)
    if not pid:
        raise Bad("project", "مشروع غير موجود: " + proj)
    code = _s(r.get("code"), "code", True, 40)
    typ = _s(r.get("type"), "type", True, 60)
    use = _s(r.get("use"), "use", maxlen=20) or "residential"
    use = {"سكني": "residential", "تجاري": "commercial", "مكتبي": "office", "مخزن": "warehouse", "أرض": "land"}.get(use, use)
    if use not in USES:
        raise Bad("use", "الاستخدام: " + "، ".join(USES))
    area = _n(r.get("area"), "area") or 0.0
    bays = _n(r.get("parking_bays"), "parking_bays", hi=50)
    owner = _s(r.get("owner"), "owner")
    oid = x.owners.get(owner) if owner else x.own
    if owner and not oid:
        raise Bad("owner", "مالك غير موجود: " + owner)
    key = (pid, code)
    if key in x.seen:
        raise Bad("code", "مكرر داخل الملف")
    x.seen.add(key)
    cur = x.c.execute("SELECT * FROM units WHERE project_id=? AND code=?", (pid, code)).fetchone()
    data = {"project_id": pid, "code": code, "type": typ, "area": area, "use": use, "owner_id": oid, "parking_bays": int(bays) if bays is not None else None}
    if not cur:
        return {"action": "insert", "data": data}
    if cur["source"] == "import":
        fill = {k: v for k, v in (("owner_id", oid), ("parking_bays", data["parking_bays"]), ("area", area or None)) if v is not None and cur[k] in (None, 0, "")}
        return {"action": "update", "id": cur["id"], "data": fill} if fill else {"action": "skip"}
    same = cur["type"] == typ and (not area or abs((cur["area"] or 0) - area) < 0.01) and (owner is None or cur["owner_id"] == oid) and (bays is None or cur["parking_bays"] == int(bays))
    return {"action": "skip"} if same else {"action": "conflict", "msg": "وحدة موجودة (يدوية/مبذورة) ببيانات مختلفة — لم تُعدَّل"}


def _tenant_maps(c) -> dict:
    m = {"id": {}, "cr": {}, "namephone": {}, "name": {}}
    for t in c.execute("SELECT id, name, phone, id_number, cr_number, source FROM tenants_l").fetchall():
        if t["id_number"]:
            m["id"][pii.dec(t["id_number"])] = t
        if t["cr_number"]:
            m["cr"][t["cr_number"]] = t
        m["namephone"][(t["name"].strip(), (t["phone"] or "").strip())] = t
        m["name"].setdefault(t["name"].strip(), []).append(t)
    return m


def _plan_tenants(x: Ctx, n: int, r: dict) -> dict:
    name = _s(r.get("name"), "name", True)
    if len(name) < 3:
        raise Bad("name", "أقصر من 3 أحرف")
    kind = {"شركة": "company", "فرد": "individual"}.get((_s(r.get("kind"), "kind") or "individual"), _s(r.get("kind"), "kind") or "individual")
    if kind not in ("individual", "company"):
        raise Bad("kind", "فرد أو شركة")
    ident = _s(r.get("id_number"), "id_number", maxlen=30)
    cr = _s(r.get("cr_number"), "cr_number", maxlen=30)
    phone = _s(r.get("phone"), "phone", maxlen=30)
    if phone and not set(phone) <= set("0123456789+ ()-"):
        raise Bad("phone", "أرقام فقط")
    if ident and not ident.isalnum():
        raise Bad("id_number", "حروف وأرقام فقط")
    if not (ident or cr or phone):
        raise Bad("id_number", "يلزم رقم هوية أو سجل تجاري أو هاتف للتمييز")
    key = ("id", ident) if ident else ("cr", cr) if cr else ("np", name, phone)
    if key in x.seen:
        raise Bad("name", "مكرر داخل الملف")
    x.seen.add(key)
    m = x.__dict__.setdefault("tmap", _tenant_maps(x.c))
    cur = m["id"].get(ident) if ident else m["cr"].get(cr) if cr else m["namephone"].get((name, phone or ""))
    data = {"name": name, "phone": phone or "", "id_number": ident or "", "kind": kind, "cr_number": cr or "", "email": _s(r.get("email"), "email") or "",
            "address": _s(r.get("address"), "address", maxlen=250) or "", "nationality": _s(r.get("nationality"), "nationality", maxlen=40) or ""}
    if not cur:
        return {"action": "insert", "data": data}
    if cur["source"] == "import":
        full = x.c.execute("SELECT * FROM tenants_l WHERE id=?", (cur["id"],)).fetchone()
        fill = {k: v for k, v in data.items() if k not in ("name", "id_number") and v and not (full[k] or "")}
        return {"action": "update", "id": cur["id"], "data": fill} if fill else {"action": "skip"}
    same = cur["name"].strip() == name and ((cur["phone"] or "") == (phone or "") or not phone)
    return {"action": "skip"} if same else {"action": "conflict", "msg": "مستأجر موجود (يدوي) ببيانات مختلفة — لم يُعدَّل"}


def _find_tenant(x: Ctx, ref: str):
    m = x.__dict__.setdefault("tmap", _tenant_maps(x.c))
    if ref in m["id"]:
        return m["id"][ref]
    if ref in m["cr"]:
        return m["cr"][ref]
    by = m["name"].get(ref, [])
    if len(by) > 1:
        raise Bad("tenant", "الاسم يطابق أكثر من مستأجر — استعمل رقم الهوية أو السجل")
    return by[0] if by else None


def _plan_leases(x: Ctx, n: int, r: dict) -> dict:
    code = _s(r.get("unit_code"), "unit_code", True, 40)
    units = x.c.execute("SELECT id, project_id FROM units WHERE code=? AND retained=1", (code,)).fetchall()
    if len(units) != 1:
        raise Bad("unit_code", "وحدة غير موجودة في محفظة التأجير" if not units else "الرمز يطابق أكثر من وحدة")
    ten = _find_tenant(x, _s(r.get("tenant"), "tenant", True))
    if not ten:
        raise Bad("tenant", "المستأجر غير موجود — استورد المستأجرين أولًا")
    full = x.c.execute("SELECT blacklisted FROM tenants_l WHERE id=?", (ten["id"],)).fetchone()
    if full["blacklisted"]:
        raise Bad("tenant", "المستأجر في القائمة السوداء — لا يتجاوزها الاستيراد")
    start = _d(r.get("start"), "start", True)
    months = int(_n(r.get("months"), "months", True, 1, 120))
    rents = [v for v in (_n(r.get("annual_rent"), "annual_rent", lo=0.001), _n(r.get("rent_y2"), "rent_y2", lo=0.001), _n(r.get("rent_y3"), "rent_y3", lo=0.001)) if v]
    if not rents:
        raise Bad("annual_rent", "الإيجار السنوي مطلوب")
    years = -(-months // 12)
    if len(rents) < years:
        rents = rents + [rents[-1]] * (years - len(rents))  # an unstated later year keeps the last stated rent — no invented increase
    freq = int(_n(r.get("frequency"), "frequency", lo=1, hi=12) or 12)
    if freq not in (1, 2, 3, 4, 6, 12):
        raise Bad("frequency", "يجب أن يقسم 12")
    number = _s(r.get("number"), "number", maxlen=30)
    end = addm(start, months) - dt.timedelta(days=1)
    signed = str(r.get("signed") or "").strip().lower() in YES
    cur = x.c.execute("SELECT * FROM leases WHERE number=?", (number,)).fetchone() if number else None
    cur = cur or x.c.execute("SELECT * FROM leases WHERE unit_id=? AND start=? AND status<>'terminated'", (units[0]["id"], start.isoformat())).fetchone()
    if cur:
        same = cur["unit_id"] == units[0]["id"] and cur["tenant_id"] == ten["id"] and cur["start"] == start.isoformat()
        return {"action": "skip"} if same else {"action": "conflict", "msg": "عقد موجود بالرقم أو بالوحدة وتاريخ البداية ببيانات مختلفة — لم يُعدَّل"}
    key = ("lease", units[0]["id"], start)
    if key in x.seen:
        raise Bad("start", "عقدان لوحدة واحدة ببداية واحدة داخل الملف")
    x.seen.add(key)
    if signed:
        if _overlap(x.c, units[0]["id"], start, end):
            raise Bad("start", "يتداخل مع عقد ساري على الوحدة")
        ranges = x.__dict__.setdefault("ranges", [])
        if any(uid == units[0]["id"] and s0 <= end and e0 >= start for uid, s0, e0 in ranges):
            raise Bad("start", "يتداخل مع عقد آخر في الملف على الوحدة")
        ranges.append((units[0]["id"], start, end))
    return {"action": "insert", "data": {"unit_id": units[0]["id"], "tenant_id": ten["id"], "start": start, "months": months, "rents": rents, "frequency": freq,
                                          "prep_free_days": int(_n(r.get("prep_free_days"), "prep_free_days", hi=180) or 0), "deposit": _n(r.get("deposit"), "deposit", hi=1_000_000) or 0,
                                          "number": number, "municipality_ref": _s(r.get("municipality_ref"), "municipality_ref", maxlen=60) or "", "signed": signed}}


def _plan_payments(x: Ctx, n: int, r: dict) -> dict:
    num = _s(r.get("lease_number"), "lease_number", True, 30)
    l = x.c.execute("SELECT id, status FROM leases WHERE number=?", (num,)).fetchone()
    if not l or l["status"] == "draft":
        raise Bad("lease_number", "عقد غير موجود أو مسودة")
    amount = r3(_n(r.get("amount"), "amount", True, lo=0.001))
    paid_on = _d(r.get("paid_on"), "paid_on", True)
    if paid_on > today():
        raise Bad("paid_on", "تاريخ في المستقبل")
    due_date = _d(r.get("due_date"), "due_date")
    mraw = _s(r.get("method"), "method", maxlen=20)
    method = METHOD_WORDS.get(mraw.lower() if mraw else "", None) if mraw else None
    if mraw and not method:
        raise Bad("method", "طريقة غير معروفة")
    ref = _s(r.get("reference"), "reference", maxlen=80) or ""
    # the fingerprint first: a payment this file already brought in is skipped, not reported as exceeding the (now settled) balance
    base = f"{num}|{due_date or ''}|{amount}|{paid_on}|{ref}"
    occ = x.__dict__.setdefault("occ", {})
    occ[base] = occ.get(base, 0) + 1
    key = hashlib.sha256(f"{base}|{occ[base]}".encode()).hexdigest()[:24]
    if x.c.execute("SELECT 1 FROM rent_payments WHERE import_key LIKE ?", (key + ":%",)).fetchone():
        return {"action": "skip"}
    due_id = None
    left = x.__dict__.setdefault("left", {})
    if due_date:
        d = x.c.execute("SELECT id, amount, paid FROM rent_dues WHERE lease_id=? AND due=? AND COALESCE(kind,'rent')<>'deposit'", (l["id"], due_date.isoformat())).fetchone()
        if not d:
            raise Bad("due_date", "لا استحقاق في هذا التاريخ")
        due_id = d["id"]
        room = left.setdefault(("d", due_id), r3(d["amount"] - d["paid"]))
        if amount > room + EPS:
            raise Bad("amount", f"يتجاوز المتبقي على الاستحقاق ({room:,.3f})")
        left[("d", due_id)] = r3(room - amount)
        left[("l", l["id"])] = r3(left.get(("l", l["id"]), r3(x.c.execute("SELECT COALESCE(SUM(amount-paid),0) FROM rent_dues WHERE lease_id=?", (l["id"],)).fetchone()[0])) - amount)
    else:
        room = left.setdefault(("l", l["id"]), r3(x.c.execute("SELECT COALESCE(SUM(amount-paid),0) FROM rent_dues WHERE lease_id=?", (l["id"],)).fetchone()[0]))
        if amount > room + EPS:
            raise Bad("amount", f"يتجاوز المتبقي على العقد ({room:,.3f})")
        left[("l", l["id"])] = r3(room - amount)
    return {"action": "insert", "data": {"lease_id": l["id"], "amount": amount, "paid_on": paid_on, "method": method, "ref": ref, "due_id": due_id, "key": key}}


def _plan_cheques(x: Ctx, n: int, r: dict) -> dict:
    num = _s(r.get("lease_number"), "lease_number", True, 30)
    l = x.c.execute("SELECT id, status FROM leases WHERE number=?", (num,)).fetchone()
    if not l or l["status"] == "draft":
        raise Bad("lease_number", "عقد غير موجود أو مسودة")
    cno, bank = _s(r.get("number"), "number", True, 30), _s(r.get("bank"), "bank", True, 60)
    amount = r3(_n(r.get("amount"), "amount", True, lo=0.001))
    due = _d(r.get("due_date"), "due_date", True)
    rec = _d(r.get("received_on"), "received_on") or today()
    st = {"محفوظ": "held", "مودَع": "deposited", "مودع": "deposited", "held": "held", "deposited": "deposited"}.get((_s(r.get("status"), "status") or "held").lower())
    if not st:
        raise Bad("status", "محفوظ أو مودَع (الصرف يُسجَّل من شاشة الشيكات)")
    if (cno, bank) in x.seen:
        raise Bad("number", "شيك مكرر داخل الملف")
    x.seen.add((cno, bank))
    if x.c.execute("SELECT 1 FROM cheques WHERE number=? AND bank=? AND status NOT IN ('cancelled','returned')", (cno, bank)).fetchone():
        return {"action": "skip"}
    return {"action": "insert", "data": {"lease_id": l["id"], "number": cno, "bank": bank, "amount": amount, "due": due, "rec": rec, "drawer": _s(r.get("drawer"), "drawer") or "", "status": st}}


PLANNERS = {"units": _plan_units, "tenants": _plan_tenants, "leases": _plan_leases, "payments": _plan_payments, "cheques": _plan_cheques}


def plan(c, kind: str, table: list[tuple[int, dict]]) -> dict:
    x = Ctx(c, kind)
    labels = {k: lab for k, lab, _ in SPECS[kind]}
    errors, conflicts, items = [], [], []
    for n, r in table:
        try:
            p = PLANNERS[kind](x, n, r)
        except Bad as e:
            errors.append({"row": n, "col": labels.get(e.col, e.col), "reason": e.reason})
            continue
        p["row"] = n
        if p["action"] == "conflict":
            conflicts.append({"row": n, "reason": p["msg"]})
        items.append(p)
    count = {a: sum(1 for i in items if i["action"] == a) for a in ("insert", "update", "skip", "conflict")}
    return {"kind": kind, "rows": len(table), "errors": errors, "conflicts": conflicts, "counts": count, "items": items}


# ----------------------------------------------------------------------------------------------------- applying (one transaction)
def apply(c, kind: str, items: list[dict], batch: int, u: dict) -> None:
    for it in items:
        d = it.get("data") or {}
        if it["action"] == "insert":
            if kind == "units":
                c.execute("INSERT INTO units(project_id,code,type,area,status,retained,use,owner_id,parking_bays,source,import_batch) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                          (d["project_id"], d["code"], d["type"], d["area"], "a", 1, d["use"], d["owner_id"], d["parking_bays"], "import", str(batch)))
            elif kind == "tenants":
                t = TenantIn(name=d["name"], phone=d["phone"], id_number=d["id_number"], kind=d["kind"], cr_number=d["cr_number"], email=d["email"], address=d["address"],
                             nationality=d["nationality"])
                tid = c.execute(f"INSERT INTO tenants_l({TENANT_SQL_COLS}) VALUES({','.join('?' * 16)})", _tenant_cols(t)).lastrowid
                c.execute("UPDATE tenants_l SET source='import', import_batch=? WHERE id=?", (str(batch), tid))
            elif kind == "leases":
                lid = create_lease(c, LeaseNew(unit_id=d["unit_id"], tenant_id=d["tenant_id"], start=d["start"], months=d["months"], rents=d["rents"], frequency=d["frequency"],
                                               prep_free_days=d["prep_free_days"], deposit=d["deposit"], municipality_ref=d["municipality_ref"]), u, "import",
                                   bill_deposit=d["signed"], number=d["number"])
                if d["signed"]:
                    sign_lease(c, lid, u, signed_at=d["start"])
                    c.execute("UPDATE leases SET status='expired' WHERE id=? AND \"end\"<?", (lid, today().isoformat()))
            elif kind == "payments":
                apply_payment(c, d["lease_id"], d["amount"], d["paid_on"], d["method"], d["ref"] or None, u["name"], d["due_id"], None, "استيراد تاريخي", import_key=d["key"], receipted=False)
            elif kind == "cheques":
                y = d["rec"].year
                no = f"CQ-{y}-{next_no(c, f'chqrcpt:{y}'):05d}"
                c.execute("""INSERT INTO cheques(direction,lease_id,number,bank,amount,due_date,received_on,status,drawer,receipt_no,created_by,deposited_on) VALUES('in',?,?,?,?,?,?,?,?,?,?,?)""",
                          (d["lease_id"], d["number"], d["bank"], d["amount"], d["due"].isoformat(), d["rec"].isoformat(), d["status"], d["drawer"], no, u["name"],
                           today().isoformat() if d["status"] == "deposited" else None))
        elif it["action"] == "update" and d:
            table = {"units": "units", "tenants": "tenants_l"}[kind]
            sets = ", ".join(f"{k}=?" for k in d)
            c.execute(f"UPDATE {table} SET {sets} WHERE id=?", (*d.values(), it["id"]))


# ----------------------------------------------------------------------------------------------------- endpoints
def _public(rep: dict, **extra) -> dict:
    out = {k: v for k, v in rep.items() if k != "items"}
    out["preview"] = [{"row": i["row"], "action": i["action"]} for i in rep["items"][:50]]
    return out | extra


@router.post("/api/lea/import/{kind}")
async def api_import(kind: str, file: UploadFile = File(...), confirm: str = Form(""), u=Depends(act_as("leasing_approve"))):
    """Without `confirm`: a dry-run. With `confirm` = the SHA-256 the dry-run returned: the same file is checked again and, if clean, applied."""
    if kind not in SPECS:
        raise HTTPException(404, "نوع استيراد غير معروف")
    data = await file.read(MAX_BYTES + 1)
    sha = hashlib.sha256(data).hexdigest()
    table = read_table(file.filename or "", data, kind)
    c = db()
    rep = plan(c, kind, table)
    base = {"filename": file.filename, "sha256": sha}
    if not confirm:
        return _public(rep, **base, committed=False, mode="dry-run")
    if confirm != sha:
        raise HTTPException(409, "الملف المرفوع يختلف عن الذي فُحص — أعد الفحص")
    prev = c.execute("SELECT id FROM import_batches WHERE kind=? AND sha256=?", (kind, sha)).fetchone()
    if prev:
        raise HTTPException(409, f"استُورد هذا الملف سابقًا (الدفعة {prev['id']})")
    if rep["errors"]:
        return _public(rep, **base, committed=False, mode="blocked")
    batch = c.execute("INSERT INTO import_batches(kind,filename,sha256,rows_ok,rows_err,rows_conflict,by_user,at) VALUES(?,?,?,?,?,?,?,?)",
                      (kind, file.filename, sha, rep["counts"]["insert"] + rep["counts"]["update"], 0, rep["counts"]["conflict"], u["name"], now_s())).lastrowid
    apply(c, kind, rep["items"], batch, u)
    audit(c, "استيراد " + KIND_LABEL[kind], f"{file.filename} · دفعة {batch} · جديد {rep['counts']['insert']} · محدَّث {rep['counts']['update']} · متجاوَز {rep['counts']['skip']} · تعارض {rep['counts']['conflict']}")
    c.commit()
    return _public(rep, **base, committed=True, mode="commit", batch=batch)


@router.get("/api/lea/import/batches")
def api_import_batches(_=Depends(act_as("leasing_approve"))):
    return rows(db().execute("SELECT * FROM import_batches ORDER BY id DESC LIMIT 100"))


@router.get("/api/lea/import/{kind}/template.xlsx")
def api_import_template(kind: str, _=Depends(act_as("leasing_approve"))):
    if kind not in SPECS:
        raise HTTPException(404, "نوع استيراد غير معروف")
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "البيانات"
    ws.sheet_view.rightToLeft = True
    ws.append([lab for _k, lab, _r in SPECS[kind]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i in range(1, len(SPECS[kind]) + 1):
        ws.column_dimensions[chr(64 + i)].width = 22
    hs = wb.create_sheet("تعليمات")
    hs.sheet_view.rightToLeft = True
    hs.append([f"قالب استيراد {KIND_LABEL[kind]} — الصفّ الأول عناوين الأعمدة، ابدأ بياناتك من الصف الثاني في ورقة «البيانات»."])
    hs.append(["الأعمدة المطلوبة: " + "، ".join(lab for _k, lab, rq in SPECS[kind] if rq)])
    hs.append(["التواريخ بصيغة 2026-01-31. لا تستعمل صيغًا أو معادلات (تُرفض)؛ الصق القيم فقط. الحد 5000 صف و5 ميغابايت."])
    hs.append(["الترتيب: الوحدات ثم المستأجرون ثم العقود ثم المدفوعات والشيكات. الفحص المسبق لا يكتب شيئًا؛ والتنفيذ بعد التأكيد فقط."])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=import-{kind}.xlsx", "Cache-Control": "no-store"})


__all__ = ["router"]
