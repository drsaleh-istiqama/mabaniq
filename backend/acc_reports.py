"""Mabaniq 0.24.0 — the four financial statements, their exports, the year-end close and the ERPNext export (design §4.7, phase 3b).

Every statement is computed from the posted ledger and carries its own proof: the balance sheet states Assets = Liabilities + Equity, the cash-flow statement
reconciles to the real change in the cash accounts, and the equity statement reconciles to the balance sheet — each is reported, not assumed. A closed year's
profit is carried to retained earnings by ONE closing journal; the income / cash-flow / equity statements ignore that journal (so a closed year still shows its profit),
and the balance sheet includes it (so the profit sits where it belongs).
"""
from __future__ import annotations

import csv
import datetime as dt
import io
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, Field

from . import ledger, pdfgen
from .acc_ap import _active
from .auth import need
from .common import act_as, audit, db, now_s, rows, today

router = APIRouter()
r3 = ledger.r3
EPS = ledger.EPS
POSTED = "j.status IN ('posted','reversed')"
NOT_CLOSE = "COALESCE(j.source_kind,'')<>'yearclose'"


# ----------------------------------------------------------------------------------------------------- balances
def balances(c, frm: str | None, to: str, include_close: bool, cc: int | None = None) -> dict[int, float]:
    """account id → Σ(debit − credit) of posted journals dated frm..to (frm None = from the beginning)."""
    q = f"SELECT l.account_id, SUM(l.debit - l.credit) b FROM acc_lines l JOIN acc_journals j ON j.id=l.journal_id WHERE {POSTED} AND j.jdate<=?"
    args: list = [to]
    if frm:
        q += " AND j.jdate>=?"
        args.append(frm)
    if not include_close:
        q += f" AND {NOT_CLOSE}"
    if cc:
        q += " AND l.cost_center_id=?"
        args.append(cc)
    return {r["account_id"]: r3(r["b"]) for r in c.execute(q + " GROUP BY l.account_id", args).fetchall()}


def _accounts(c) -> list[dict]:
    return rows(c.execute("SELECT * FROM acc_accounts ORDER BY code"))


def _tree(accs: list[dict], amount: dict[int, float], types: tuple, sign: int) -> tuple[list[dict], float]:
    """Rows (code, name, level, amount) of the accounts of `types`, leaves with a balance + their headers rolled up; `sign` +1 shows debit-positive, −1 credit-positive."""
    by_id = {a["id"]: a for a in accs}
    tot: dict[int, float] = {}
    for a in accs:
        if a["postable"] and a["type"] in types and abs(amount.get(a["id"], 0)) > EPS:
            v = amount[a["id"]] * sign
            p = a["id"]
            while p:
                tot[p] = r3(tot.get(p, 0) + v)
                p = (by_id[p]["parent_id"]) or 0
    out = []
    for a in accs:
        if a["id"] in tot:
            lvl, p = 0, a["parent_id"]
            while p:
                lvl += 1
                p = by_id[p]["parent_id"]
            out.append({"code": a["code"], "name": a["name"], "level": lvl, "amount": tot[a["id"]], "header": not a["postable"]})
    return out, r3(sum(v for k, v in tot.items() if by_id[k]["postable"]))


def _flat(sections: list[tuple[str, list[dict], float]]) -> list[dict]:
    out = []
    for title, items, total in sections:
        out.append({"cells": [title, ""], "bold": True, "level": 0})
        out += [{"cells": [f"{x['code']}  {x['name']}", x["amount"]], "bold": x["header"], "level": x["level"] + 1} for x in items]
        out.append({"cells": [f"إجمالي {title}", total], "bold": True, "level": 0})
    return out


# ----------------------------------------------------------------------------------------------------- the statements
def balance_sheet(c, to: str, cc: int | None = None) -> dict:
    accs = _accounts(c)
    bal = balances(c, None, to, True, cc)
    assets, ta = _tree(accs, bal, ("asset",), 1)
    liab, tl = _tree(accs, bal, ("liability",), -1)
    eq, te = _tree(accs, bal, ("equity",), -1)
    earnings = r3(-sum(v for a in accs if a["type"] in ("revenue", "expense") for v in [bal.get(a["id"], 0)]))  # unclosed profit of the open year(s)
    total_eq = r3(te + earnings)
    secs = [("الأصول", assets, ta), ("الالتزامات", liab, tl), ("حقوق الملكية", eq + [{"code": "", "name": "أرباح (خسائر) الفترة غير المرحَّلة", "level": 0, "amount": earnings, "header": False}], total_eq)]
    diff = r3(ta - tl - total_eq)
    return {"key": "balance_sheet", "title": "قائمة المركز المالي", "as_of": to, "assets": ta, "liabilities": tl, "equity": total_eq, "earnings": earnings, "difference": diff, "balanced": abs(diff) <= EPS,
            "columns": ["البند", "المبلغ (ر.ع)"], "rows": _flat(secs) + [{"cells": ["إجمالي الالتزامات وحقوق الملكية", r3(tl + total_eq)], "bold": True, "level": 0}]}


def income_statement(c, frm: str, to: str, cc: int | None = None) -> dict:
    accs = _accounts(c)
    bal = balances(c, frm, to, False, cc)
    rev, tr = _tree(accs, bal, ("revenue",), -1)
    exp, te = _tree(accs, bal, ("expense",), 1)
    net = r3(tr - te)
    ch = c.execute(f"""SELECT COALESCE(SUM(l.credit),0) c, COALESCE(SUM(l.debit),0) d FROM acc_lines l JOIN acc_journals j ON j.id=l.journal_id JOIN acc_accounts a ON a.id=l.account_id
                       WHERE {POSTED} AND a.system_key='charity_payable' AND j.jdate>=? AND j.jdate<=?""", (frm, to)).fetchone()
    notes = [{"text": "تبرعات التأخير المحصَّلة للجهات الخيرية في الفترة — التزام لا إيراد، ولا تدخل في الأرباح أعلاه", "amount": r3(ch["c"])},
             {"text": "منها ما صُرف للجهات الخيرية في الفترة", "amount": r3(ch["d"])}]
    rows_ = _flat([("الإيرادات", rev, tr), ("المصروفات", exp, te)]) + [{"cells": ["صافي الربح (الخسارة)", net], "bold": True, "level": 0}]
    return {"key": "income_statement", "title": "قائمة الدخل", "from": frm, "to": to, "revenue": tr, "expenses": te, "net": net, "notes": notes,
            "columns": ["البند", "المبلغ (ر.ع)"], "rows": rows_ + [{"cells": [n["text"], n["amount"]], "bold": False, "level": 1} for n in notes]}


def cash_flow(c, frm: str, to: str) -> dict:
    """Indirect method. Every non-cash account's change in the period is a cash effect (−Δ of its debit-balance); revenue and expense accounts together are the net profit.
    Accounts are classed operating (default) / investing / financing by `cash_class`; the result is reconciled to the real change of the cash accounts."""
    accs = _accounts(c)
    cash_ids = {a["id"] for a in accs if a["cash_class"] == "cash"}
    day0 = (dt.date.fromisoformat(frm) - dt.timedelta(days=1)).isoformat()
    op, cl = balances(c, None, day0, False), balances(c, None, to, False)
    delta = {a["id"]: r3(cl.get(a["id"], 0) - op.get(a["id"], 0)) for a in accs}
    net = r3(-sum(delta[a["id"]] for a in accs if a["type"] in ("revenue", "expense")))
    groups: dict[str, list[dict]] = {"operating": [], "investing": [], "financing": []}
    for a in accs:
        if a["id"] in cash_ids or a["type"] in ("revenue", "expense") or abs(delta[a["id"]]) <= EPS:
            continue
        cls = a["cash_class"] if a["cash_class"] in ("investing", "financing") else "operating"
        groups[cls].append({"name": f"{a['code']}  {a['name']}", "amount": r3(-delta[a["id"]])})
    ops = [{"name": "صافي الربح (الخسارة)", "amount": net}] + groups["operating"]
    totals = {"operating": r3(sum(x["amount"] for x in ops)), "investing": r3(sum(x["amount"] for x in groups["investing"])), "financing": r3(sum(x["amount"] for x in groups["financing"]))}
    change = r3(sum(totals.values()))
    cash_open, cash_close = r3(sum(op.get(i, 0) for i in cash_ids)), r3(sum(cl.get(i, 0) for i in cash_ids))
    actual = r3(cash_close - cash_open)
    out_rows = []
    for title, items, key in (("التدفقات النقدية من الأنشطة التشغيلية", ops, "operating"), ("التدفقات النقدية من الأنشطة الاستثمارية", groups["investing"], "investing"),
                              ("التدفقات النقدية من الأنشطة التمويلية", groups["financing"], "financing")):
        out_rows.append({"cells": [title, ""], "bold": True, "level": 0})
        out_rows += [{"cells": [x["name"], x["amount"]], "bold": False, "level": 1} for x in items]
        out_rows.append({"cells": [f"صافي {title.split(' من ')[1]}", totals[key]], "bold": True, "level": 0})
    out_rows += [{"cells": ["صافي التغير في النقد", change], "bold": True, "level": 0}, {"cells": ["النقد في بداية الفترة", cash_open], "bold": False, "level": 0},
                 {"cells": ["النقد في نهاية الفترة", cash_close], "bold": True, "level": 0}]
    return {"key": "cash_flow", "title": "قائمة التدفقات النقدية (غير مباشرة)", "from": frm, "to": to, "net_profit": net, **totals, "change": change, "cash_open": cash_open, "cash_close": cash_close,
            "difference": r3(change - actual), "reconciled": abs(change - actual) <= EPS, "columns": ["البند", "المبلغ (ر.ع)"], "rows": out_rows}


def equity_changes(c, frm: str, to: str) -> dict:
    accs = _accounts(c)
    day0 = (dt.date.fromisoformat(frm) - dt.timedelta(days=1)).isoformat()
    eq_ids = {a["id"] for a in accs if a["type"] == "equity"}
    pl_ids = {a["id"] for a in accs if a["type"] in ("revenue", "expense")}
    open_all = balances(c, None, day0, True)
    opening = r3(-sum(v for k, v in open_all.items() if k in eq_ids or k in pl_ids))  # equity accounts + not-yet-closed profit
    move = balances(c, frm, to, False)
    net = r3(-sum(v for k, v in move.items() if k in pl_ids))
    other = r3(-sum(v for k, v in move.items() if k in eq_ids))
    closing = r3(opening + net + other)
    bs = balance_sheet(c, to)
    lines = [{"cells": ["حقوق الملكية في بداية الفترة", opening], "bold": True, "level": 0}, {"cells": ["صافي ربح (خسارة) الفترة", net], "bold": False, "level": 1},
             {"cells": ["حركات حقوق الملكية الأخرى (رأس المال، الاحتياطيات، التوزيعات)", other], "bold": False, "level": 1}, {"cells": ["حقوق الملكية في نهاية الفترة", closing], "bold": True, "level": 0}]
    return {"key": "equity_changes", "title": "قائمة التغير في حقوق الملكية", "from": frm, "to": to, "opening": opening, "net_profit": net, "other": other, "closing": closing,
            "balance_sheet_equity": bs["equity"], "reconciled": abs(closing - bs["equity"]) <= EPS, "columns": ["البند", "المبلغ (ر.ع)"], "rows": lines}


def _range(frm: str | None, to: str | None) -> tuple[str, str]:
    t = dt.date.fromisoformat(to) if to else today()
    f = dt.date.fromisoformat(frm) if frm else dt.date(t.year, 1, 1)
    if f > t:
        raise HTTPException(400, "بداية المدة بعد نهايتها")
    return f.isoformat(), t.isoformat()


STATEMENTS = ("balance_sheet", "income_statement", "cash_flow", "equity_changes")


def build(c, kind: str, frm: str | None, to: str | None, cc: int | None = None) -> dict:
    if kind not in STATEMENTS:
        raise HTTPException(404, "قائمة غير معروفة")
    f, t = _range(frm, to)
    if kind == "balance_sheet":
        return balance_sheet(c, t, cc)
    if kind == "income_statement":
        return income_statement(c, f, t, cc)
    if kind == "cash_flow":
        return cash_flow(c, f, t)
    return equity_changes(c, f, t)


# ----------------------------------------------------------------------------------------------------- exports
def _fmt(v) -> str:
    return f"{v:,.3f}" if isinstance(v, (int, float)) else str(v)


def _csv(stmt: dict) -> Response:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([stmt["title"], stmt.get("as_of") or f"{stmt.get('from')} → {stmt.get('to')}"])
    w.writerow(stmt["columns"])
    for r in stmt["rows"]:
        w.writerow([("   " * r["level"]) + str(r["cells"][0]), r["cells"][1]])
    return Response("﻿" + buf.getvalue(), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{stmt["key"]}.csv"'})


def _xlsx(stmt: dict) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = stmt["title"][:30]
    ws.sheet_view.rightToLeft = True
    ws.append([stmt["title"], stmt.get("as_of") or f"{stmt.get('from')} → {stmt.get('to')}"])
    ws.append(stmt["columns"])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for r in stmt["rows"]:
        ws.append([("   " * r["level"]) + str(r["cells"][0]), r["cells"][1] if r["cells"][1] != "" else None])
        if r["bold"]:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.cell(ws.max_row, 2).font = Font(bold=True)
        ws.cell(ws.max_row, 2).number_format = "#,##0.000"
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 18
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{stmt["key"]}.xlsx"'})


def _pdf(c, stmt: dict) -> Response:
    from .paperwork import brand
    p = pdfgen.Doc(brand(c), "financial_statements", stmt["key"], "", "")
    p.add_page()
    p.heading(stmt["title"], (f"كما في {stmt['as_of']} · {pdfgen.hijri(stmt['as_of'])}" if stmt.get("as_of") else f"من {stmt['from']} إلى {stmt['to']} · {pdfgen.hijri(stmt['to'])}"))
    p.grid(stmt["columns"], [[("      " * r["level"]) + str(r["cells"][0]), _fmt(r["cells"][1]) if r["cells"][1] != "" else ""] for r in stmt["rows"]], widths=(7, 2))
    proof = {"balance_sheet": "الأصول = الالتزامات + حقوق الملكية" if stmt.get("balanced") else f"تنبيه: الفرق {stmt.get('difference')}",
             "cash_flow": "طابقت التدفقات الفرق الفعلي في حسابات النقد" if stmt.get("reconciled") else f"تنبيه: الفرق {stmt.get('difference')}",
             "equity_changes": "طابقت حقوق الملكية قائمة المركز المالي" if stmt.get("reconciled") else "تنبيه: لا تطابق قائمة المركز المالي"}.get(stmt["key"])
    if proof:
        p.para(proof, 8.5, muted=True)
    p.para("أُعدّت من الدفاتر المرحَّلة في منصة مبانيك. تبرعات التأخير التزامات للجهات الخيرية لا إيرادات.", 8.5, muted=True)
    return Response(bytes(p.output()), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{stmt["key"]}.pdf"', "Cache-Control": "no-store"})


@router.get("/api/acc/statements/{kind}")
def api_statement(kind: str, frm: str | None = None, to: str | None = None, cc: int | None = None, fmt: Literal["json", "csv", "xlsx", "pdf"] = "json", _=Depends(need("accounting"))):
    c = db()
    stmt = build(c, kind, frm, to, cc)
    if fmt == "csv":
        return _csv(stmt)
    if fmt == "xlsx":
        return _xlsx(stmt)
    if fmt == "pdf":
        return _pdf(c, stmt)
    return stmt


# ----------------------------------------------------------------------------------------------------- year-end close
class CloseIn(BaseModel):
    year: int = Field(ge=2000, le=2100)
    ignore_pending: bool = False


@router.post("/api/acc/year-close")
def api_year_close(b: CloseIn, u=Depends(act_as("accounting_approve"))):
    """Carry the year's profit to retained earnings with one journal dated 31 December, then close every month of the year."""
    c = db()
    start = _active(c)
    if b.year >= today().year:
        raise HTTPException(409, "لا تُقفل سنة لم تنتهِ بعد")
    if f"{b.year}-12-31" < start:
        raise HTTPException(409, "السنة تسبق بدء المحاسبة")
    if c.execute("SELECT 1 FROM acc_journals WHERE source_kind='yearclose' AND source_ref=? AND kind='auto' AND status='posted'", (str(b.year),)).fetchone():
        raise HTTPException(409, "أُقفلت هذه السنة سابقًا")
    tries = c.execute("SELECT COUNT(*) FROM acc_journals WHERE source_kind='yearclose' AND source_ref=? AND kind='auto'", (str(b.year),)).fetchone()[0]   # a reopened year can be closed again
    if c.execute("SELECT 1 FROM acc_journals WHERE status IN ('draft','pending') AND jdate>=? AND jdate<=?", (f"{b.year}-01-01", f"{b.year}-12-31")).fetchone():
        raise HTTPException(409, "في السنة قيود مسودة أو بانتظار الاعتماد — اعتمدها أو احذفها أولًا")
    if not b.ignore_pending and sum(ledger.autopost(c, dry=True)["pending"].values()):
        raise HTTPException(409, "أحداث تأجير لم تُرحَّل بعد — رحِّلها من «ترحيل الآن» (أو أقفل مع ignore_pending إن كانت معروفة)")
    accs = [a for a in _accounts(c) if a["type"] in ("revenue", "expense") and a["postable"]]
    bal = balances(c, f"{b.year}-01-01", f"{b.year}-12-31", False)
    lines, net = [], 0.0
    for a in accs:
        v = bal.get(a["id"], 0)
        if abs(v) <= EPS:
            continue
        lines.append(dict(account=a["id"], credit=v) if v > 0 else dict(account=a["id"], debit=-v))
        net = r3(net - v)
    jid = None
    if lines:
        lines.append(dict(account="retained_earnings", credit=net) if net > 0 else dict(account="retained_earnings", debit=-net))
        try:
            jid = ledger.post(c, date=dt.date(b.year, 12, 31), lines=lines, memo=f"إقفال السنة المالية {b.year} — ترحيل صافي الربح إلى الأرباح المرحّلة", source_kind="yearclose",
                              source_ref=str(b.year), source_key=f"yearclose:{b.year}" + (f":{tries + 1}" if tries else ""), by=u["name"], kind="auto")
        except ledger.LedgerError as e:
            raise HTTPException(409, str(e)) from None
    for m in range(1, 13):
        p = ledger.period_for(c, dt.date(b.year, m, 1))
        if p["status"] != "closed":
            c.execute("UPDATE acc_periods SET status='closed', closed_by=?, closed_at=? WHERE id=?", (u["name"], now_s(), p["id"]))
    audit(c, "إقفال سنة مالية", f"{b.year} · صافي الربح المرحَّل {net:,.3f} ر.ع")
    c.commit()
    return {"journal_id": jid, "net_profit": net}


class ReopenYear(BaseModel):
    reason: str = Field(min_length=8, max_length=300)


@router.post("/api/acc/year-close/{year}/reopen")
def api_year_reopen(year: int, b: ReopenYear, u=Depends(act_as("accounting_approve"))):
    """Undo a year-end close (reasoned and audited): the months reopen and the closing journal is reversed."""
    c = db()
    j = c.execute("SELECT id, status FROM acc_journals WHERE source_kind='yearclose' AND source_ref=? AND kind='auto' ORDER BY id DESC LIMIT 1", (str(year),)).fetchone()
    if not j and not c.execute("SELECT 1 FROM acc_periods WHERE year=? AND status='closed'", (year,)).fetchone():
        raise HTTPException(409, "السنة غير مقفلة")
    c.execute("UPDATE acc_periods SET status='open', reopened_by=?, reopened_at=?, reopen_reason=? WHERE year=?", (u["name"], now_s(), b.reason, year))
    rid = None
    if j and j["status"] == "posted":
        rid = ledger.reverse(c, j["id"], dt.date(year, 12, 31), f"إلغاء إقفال السنة {year}: {b.reason}", u["name"])
    audit(c, "إعادة فتح سنة مالية", f"{year} · {b.reason}")
    c.commit()
    return {"reversal_id": rid}


# ----------------------------------------------------------------------------------------------------- ERPNext / Zimam export
ROOT = {"asset": "Asset", "liability": "Liability", "equity": "Equity", "revenue": "Income", "expense": "Expense"}
PARTY = {"tenant": "Customer", "customer": "Customer", "vendor": "Supplier", "owner": "Supplier"}


@router.get("/api/acc/export/erpnext")
def api_export_erpnext(frm: str | None = None, to: str | None = None, _=Depends(need("accounting"))):
    """The chart of accounts and the posted journals in ERPNext «Journal Entry» field order — a file to import, not a live integration."""
    c = db()
    f, t = _range(frm, to)
    accs = _accounts(c)
    by_id = {a["id"]: a for a in accs}
    chart = [{"account_number": a["code"], "account_name": a["name"], "parent_account": by_id[a["parent_id"]]["name"] if a["parent_id"] else "", "is_group": 0 if a["postable"] else 1,
              "root_type": ROOT[a["type"]], "account_type": ""} for a in accs]
    entries = []
    for j in rows(c.execute(f"SELECT * FROM acc_journals j WHERE {POSTED} AND j.jdate>=? AND j.jdate<=? ORDER BY j.jdate, j.id", (f, t))):
        ls = rows(c.execute("""SELECT l.*, a.code, a.name account_name, cc.name cc_name FROM acc_lines l JOIN acc_accounts a ON a.id=l.account_id LEFT JOIN acc_cost_centers cc ON cc.id=l.cost_center_id
                               WHERE l.journal_id=? ORDER BY l.line_no""", (j["id"],)))
        entries.append({"voucher_type": "Journal Entry", "name": j["number"], "posting_date": j["jdate"], "user_remark": j["memo"], "total_debit": j["total"], "total_credit": j["total"],
                        "accounts": [{"account": x["account_name"], "account_number": x["code"], "debit_in_account_currency": x["debit"], "credit_in_account_currency": x["credit"],
                                      "cost_center": x["cc_name"] or "", "party_type": PARTY.get(x["party_type"] or "", ""), "party": str(x["party_id"] or "") if x["party_type"] in PARTY else "",
                                      "user_remark": x["memo"] or ""} for x in ls]})
    return JSONResponse({"company_currency": "OMR", "from": f, "to": t, "chart_of_accounts": chart, "journal_entries": entries},
                        headers={"Content-Disposition": 'attachment; filename="mabaniq-erpnext-export.json"'})

